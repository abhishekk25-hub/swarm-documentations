#!/usr/bin/env python3
"""Deterministic executable verifier for the clinical-trial risk-based monitoring task.

The deliverable is a quarterly risk-based monitoring (RBM) packet with two parts:
  * /logs/agent/output.xlsx        -- the structured analysis (per-trial risk
                                      scores/tiers/hours, the monitoring schedule,
                                      the overall-vs-site status conflict log, the
                                      out-of-scope exclusion log, CRA capacity
                                      allocation, the source register, the audit
                                      log and a summary)
  * /logs/agent/monitoring_plan.md -- the long-form, evidence-backed written plan
                                      (RBM methodology, risk tiering, per-trial
                                      monitoring proofs, schedule/capacity, window-
                                      and capacity-infeasibility analysis, data
                                      reconciliation, exclusions and citations)

Scoring is a transparent WEIGHTED RUBRIC. Each category is scored as its own
earned/possible ratio with smooth partial credit, then combined with fixed,
mode-agnostic weights. There is no per-mode branch, cap, floor or multiplier.

The numeric backbone -- every per-trial risk score, tier, monitoring-hour figure,
the in-scope/exclusion determination, the overall-vs-site status reconciliation and
the reference monitoring schedule -- is RECOMPUTED from the gold per-trial fields in
tests/reference_inputs/trials_reference.json (extracted verbatim from the real
ClinicalTrials.gov records) using the fully-specified rulebook constants. There is
no pre-authored answer file. The browsing check compares the agent's submitted
citations against acceptable source domains and load-bearing fact tokens (the
verifier never accesses the network). The writing checks are coverage/consistency
based (a quantified per-trial monitoring proof, every trial carried to a decision,
required sections present and internally consistent with the recomputed schedule).
"""

import argparse
import json
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REF_DIR = os.path.join(HERE, "reference_inputs")

DATE_RE = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
URL_RE = re.compile(r"https?://[^\s)\]]+", re.IGNORECASE)

# Mode-agnostic rubric weights (sum to 1.0), assigned by each deliverable's
# prominence in the work brief. The analytical workbook -- per-trial risk scoring,
# tiering, monitoring hours, the schedule, exclusions and the overall-vs-site status
# reconciliation -- is the largest deliverable and is split into numeric_backbone
# and status_reconciliation. Confirming and citing the twelve authorities
# (source_register) and the long evidence-backed monitoring plan (trial_proofs +
# report_structure) are the other primary deliverables the brief calls for, and
# audit_coverage captures carrying every trial through to a decision. Weights track
# instruction emphasis only; there is no per-mode branch, cap, floor or multiplier.
WEIGHTS = {
    "numeric_backbone": 0.26,
    "status_reconciliation": 0.12,
    "source_register": 0.22,
    "trial_proofs": 0.18,
    "audit_coverage": 0.10,
    "report_structure": 0.12,
}

REQUIRED_SHEETS = [
    "trial_risk_scores", "monitoring_schedule", "conflict_log", "exclusion_log",
    "capacity_allocation", "source_register", "audit_log", "summary",
]

REPORT_SECTIONS = [
    ["executive summary"],
    ["background", "portfolio overview", "portfolio"],
    ["methodology", "rulebook", "risk formula", "risk score"],
    ["risk tier", "tiering", "high tier", "high-risk"],
    ["regulated", "data integrity", "data-integrity", "ind", "ide"],
    ["trial-by-trial", "trial by trial", "per-trial", "monitoring proof"],
    ["monitoring schedule", "capacity allocation", "cra capacity", "staffing"],
    ["window", "infeasib", "deferr", "capacity"],
    ["reconciliation", "conflict", "overall status", "site status", "governing"],
    ["exclud", "out-of-scope", "out of scope", "data limitation", "withdrawn", "terminated"],
    ["recommendation", "mitigation"],
    ["regulatory basis", "citations", "authorities"],
]

KEY_DISCLOSURES = ["risk-based monitoring", "risk tier", "visit window", "hosting"]
# (KEY_DISCLOSURES kept domain-specific below)
KEY_DISCLOSURES = ["risk", "monitoring", "tier", "schedule"]
DECISION_KEYWORDS = [
    "schedul", "monitor", "visit", "defer", "window", "infeasib", "exclud", "withdraw",
    "terminat", "complet", "capacity", "tier", "high", "medium", "low", "decision",
    "risk", "week", "blocker",
]
BANNED_TOKENS = ("lorem", "[insert", "tbd", "todo", "placeholder", "xxxx", "<fill")

JUSTIFY_MIN_CHARS = 160
REPORT_MIN_CHARS = 9000


def norm(v):
    return "" if v is None else str(v).strip()


def num(v):
    s = norm(v).replace(",", "").replace("$", "").replace("%", "")
    s = re.sub(r"(?i)\b(kw|hours|hrs|h|participants|days|sites|weeks)\b", "", s).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        m = re.search(r"-?\d+(?:\.\d+)?", s)
        return float(m.group()) if m else None


def yesno(v):
    s = norm(v).lower()
    if s in ("yes", "y", "true", "1", "in_scope", "in scope", "regulated", "divergent", "conflict"):
        return True
    if s in ("no", "n", "false", "0", "out_of_scope", "out of scope", "none", "not regulated",
             "excluded", "not divergent", "ok", "consistent"):
        return False
    return None


def close_abs(a, b, tol):
    return a is not None and abs(a - b) <= tol


# ---------------------------------------------------------------------------
# Recompute the numeric reference from the gold fields + rulebook constants.
# ---------------------------------------------------------------------------
def _tier_of(score, rb):
    if score >= rb["high_tier_min"]:
        return "high"
    if score >= rb["medium_tier_min"]:
        return "medium"
    return "low"


def _phase_pts(phases, rb):
    pp = rb["phase_pts"]
    if not phases:
        return pp.get("NA", pp.get("", 0))
    return max(pp.get(p, pp.get("NA", 0)) for p in phases)


def _tier_from_steps(value, steps):
    """steps = [(threshold, pts), ...] descending; return pts for value >= threshold."""
    for k, v in steps:
        if value >= k:
            return v
    return 0


def build_reference():
    rb = json.load(open(os.path.join(REF_DIR, "rulebook_constants.json")))
    data = json.load(open(os.path.join(REF_DIR, "trials_reference.json")))
    trials = data["trials"]

    cap = rb["weekly_cra_hours_cap"]
    horizon = rb["horizon_weeks"]
    in_scope_set = set(rb["in_scope_statuses"])
    excl = rb["exclusion_reasons"]
    window_weeks = rb["window_weeks"]

    recs = {}
    conflicts = []
    for t in trials:
        nct = t["nct_id"]
        status = (t["overall_status"] or "").upper()
        in_scope = status in in_scope_set
        exclusion_reason = "" if in_scope else excl.get(status, "out_of_scope")

        enrollment = t["enrollment_count"] or 0
        staleness = t["staleness_days"] or 0
        regulated = bool(t["regulated"])
        estimated = (t["enrollment_basis"] or "").upper() == "ESTIMATED"

        if t["status_divergence"]:
            conflicts.append(dict(nct_id=nct, field="overall_status_vs_site_status",
                                  governing=status,
                                  site_values=",".join(t["site_statuses"])))

        rec = dict(
            nct_id=nct, overall_status=status, study_type=(t["study_type"] or "").upper(),
            phases=[p.upper() for p in (t["phases"] or [])],
            enrollment=enrollment, enrollment_basis=(t["enrollment_basis"] or "").upper(),
            num_sites=t["num_sites"] or 0, regulated=regulated,
            staleness=staleness, status_divergence=bool(t["status_divergence"]),
            in_scope=in_scope, exclusion_reason=exclusion_reason,
        )

        if in_scope:
            score = (_phase_pts(rec["phases"], rb)
                     + rb["study_type_pts"].get(rec["study_type"], 0)
                     + _tier_from_steps(enrollment, rb["enrollment_pts"])
                     + _tier_from_steps(rec["num_sites"], rb["sites_pts"])
                     + _tier_from_steps(staleness, rb["staleness_pts"])
                     + (rb["estimated_enrollment_pts"] if estimated else 0)
                     + (rb["regulated_pts"] if regulated else 0))
            tier = _tier_of(score, rb)
            hours = (rb["base_monitor_hours"] + rb["tier_monitor_hours"][tier]
                     + _tier_from_steps(rec["num_sites"], rb["sites_monitor_hours"]))
            rec.update(risk_score=score, risk_tier=tier, monitor_hours=hours,
                       window=window_weeks[tier])
        else:
            rec.update(risk_score=None, risk_tier="excluded", monitor_hours=0, window=None)
        recs[nct] = rec

    # Reference monitoring schedule (fully-specified greedy).
    active = [r for r in recs.values() if r["in_scope"]]
    order = sorted(active, key=lambda r: (-r["risk_score"], r["window"],
                                          -r["enrollment"], r["nct_id"]))
    week_hours = [0] * (horizon + 1)
    for r in order:
        hours = r["monitor_hours"]
        win = r["window"]
        assigned = None
        for w in range(1, horizon + 1):
            if w <= win and week_hours[w] + hours <= cap:
                assigned = w
                break
        if assigned is not None:
            week_hours[assigned] += hours
            r.update(decision="schedule_visit", assigned_week=assigned, blocker="")
        else:
            any_room = any(week_hours[w] + hours <= cap for w in range(1, horizon + 1))
            if any_room:
                r.update(decision="window_infeasible", assigned_week=None,
                         blocker="cannot meet visit-window cadence under CRA capacity")
            else:
                r.update(decision="defer_capacity", assigned_week=None,
                         blocker="no CRA capacity in horizon")

    active_ids = [r["nct_id"] for r in recs.values() if r["in_scope"]]
    excluded_ids = [r["nct_id"] for r in recs.values() if not r["in_scope"]]
    scheduled = [r for r in active if r["decision"] == "schedule_visit"]

    def tier_count(tier):
        return sum(1 for r in active if r["risk_tier"] == tier)

    return dict(
        recs=recs, rb=rb, week_hours=week_hours,
        active_ids=active_ids, excluded_ids=excluded_ids,
        conflicts=conflicts,
        conflict_ids=sorted({c["nct_id"] for c in conflicts}),
        high_tier_count=tier_count("high"),
        medium_tier_count=tier_count("medium"),
        low_tier_count=tier_count("low"),
        regulated_count=sum(1 for r in active if r["regulated"]),
        status_conflicts_count=sum(1 for r in recs.values() if r["status_divergence"]),
        window_infeasible_count=sum(1 for r in active if r["decision"] == "window_infeasible"),
        capacity_deferred_count=sum(1 for r in active if r["decision"] == "defer_capacity"),
        scheduled_count=len(scheduled),
        high_oversight_completions=sum(1 for r in scheduled if r["risk_tier"] == "high"),
        total_monitor_hours=sum(week_hours),
    )


def load_sources():
    with open(os.path.join(REF_DIR, "sources_reference.json")) as f:
        return json.load(f)["sources"]


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------
def read_sheet(wb, name):
    if name not in wb.sheetnames:
        return None
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    header = [norm(h).lower() for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(c is None or norm(c) == "" for c in raw):
            continue
        out.append({header[i]: (raw[i] if i < len(raw) else None) for i in range(len(header))})
    return out


def kv_sheet(wb, name):
    rows = read_sheet(wb, name) or []
    out = {}
    for r in rows:
        keys = list(r.keys())
        if len(keys) >= 2:
            k = norm(r[keys[0]]).lower()
            if k:
                out[k] = r[keys[1]]
    return out


def get(row, *names):
    for n in names:
        if n in row and norm(row[n]) != "":
            return row[n]
    return None


def _nct(row):
    return norm(get(row, "nct_id", "nct_number", "nct", "id", "study_id")).upper()


# ---------------------------------------------------------------------------
# Category scorers -- each returns (earned, possible, lines [, extra])
# ---------------------------------------------------------------------------
def score_numeric_backbone(ref, wb):
    R = ref["recs"]
    active = {nct: R[nct] for nct in ref["active_ids"]}
    scores = read_sheet(wb, "trial_risk_scores") or []
    by_id = {}
    for r in scores:
        nid = _nct(r)
        if nid and nid not in by_id:
            by_id[nid] = r
    sched = read_sheet(wb, "monitoring_schedule") or []
    sched_by_id = {}
    for r in sched:
        nid = _nct(r)
        if nid and nid not in sched_by_id:
            sched_by_id[nid] = r
    exclusions = read_sheet(wb, "exclusion_log") or []
    rollup = kv_sheet(wb, "summary")

    e = p = 0
    lines = []

    def per_active(label, fn):
        nonlocal e, p
        earned = sum(1 for nid in active if nid in by_id and fn(nid, by_id[nid]))
        e += earned
        p += len(active)
        lines.append(f"    {earned:>3}/{len(active):<3} trial: {label}")

    per_active("risk score recomputed", lambda nid, row: close_abs(
        num(get(row, "risk_score", "score")), active[nid]["risk_score"], 0.5))
    per_active("risk tier correct", lambda nid, row: norm(
        get(row, "risk_tier", "tier")).lower() == active[nid]["risk_tier"])
    per_active("monitoring hours correct", lambda nid, row: close_abs(
        num(get(row, "monitoring_hours", "monitor_hours", "cra_hours", "hours")),
        active[nid]["monitor_hours"], 0.5))

    # scheduled decision categorical match
    def decision_match(nid, row):
        want = active[nid]["decision"]
        got = norm(get(row, "monitoring_decision", "scheduled_decision", "decision", "disposition")).lower()
        if not got:
            return False
        if want == "schedule_visit":
            return "sched" in got or "visit" in got or "monitor" in got
        if want == "window_infeasible":
            return "window" in got or "infeasib" in got or "cadence" in got
        if want == "defer_capacity":
            return "defer" in got or "capacity" in got
        return False
    dec = 0
    for nid in active:
        row = sched_by_id.get(nid) or by_id.get(nid)
        if row and decision_match(nid, row):
            dec += 1
    e += dec
    p += len(active)
    lines.append(f"    {dec:>3}/{len(active):<3} trial: monitoring decision matches reference")

    # in-scope flag correct over ALL trials
    insc = 0
    for nid, rec in R.items():
        row = by_id.get(nid)
        if row is not None:
            v = yesno(get(row, "in_scope", "in_scope_flag", "active"))
            if v is not None and v == rec["in_scope"]:
                insc += 1
    e += insc
    p += len(R)
    lines.append(f"    {insc:>3}/{len(R):<3} trial: in-scope/out-of-scope determination correct")

    # exclusion reasons correct (per excluded trial)
    xids = set(ref["excluded_ids"])
    xseen = {_nct(x) for x in exclusions if _nct(x)}
    xreason = 0
    excl_by_id = {}
    for x in exclusions:
        nid = _nct(x)
        if nid and nid not in excl_by_id:
            excl_by_id[nid] = x
    for nid in xids:
        row = excl_by_id.get(nid)
        if row:
            got = norm(get(row, "exclusion_reason", "reason", "out_of_scope_reason")).lower()
            want = R[nid]["exclusion_reason"].lower()
            # accept the canonical token or the underlying status word
            stat = R[nid]["overall_status"].lower().replace("_", " ")
            if want and (want in got or got in want or stat.split()[0] in got):
                xreason += 1
    e += xreason
    p += len(xids)
    lines.append(f"    {xreason:>3}/{len(xids):<3} excluded trial: exclusion reason correct")

    # aggregate rollup metrics
    def rnum(*names):
        for n in names:
            if n in rollup:
                return num(rollup[n])
        return None

    roll_checks = [
        ("high_tier_count", close_abs(rnum("high_tier_count", "high_risk_count", "high_count"), ref["high_tier_count"], 0.5)),
        ("medium_tier_count", close_abs(rnum("medium_tier_count", "medium_count"), ref["medium_tier_count"], 0.5)),
        ("low_tier_count", close_abs(rnum("low_tier_count", "low_count"), ref["low_tier_count"], 0.5)),
        ("regulated_count", close_abs(rnum("regulated_count", "regulated_trials"), ref["regulated_count"], 0.5)),
        ("status_conflicts", close_abs(rnum("status_conflicts_count", "status_conflicts", "conflicts_count"), ref["status_conflicts_count"], 0.5)),
        ("window_infeasible", close_abs(rnum("window_infeasible_count", "window_infeasible"), ref["window_infeasible_count"], 0.5)),
        ("capacity_deferred", close_abs(rnum("capacity_deferred_count", "deferred_count"), ref["capacity_deferred_count"], 0.5)),
        ("scheduled_count", close_abs(rnum("scheduled_count", "monitored_count", "visits_scheduled"), ref["scheduled_count"], 0.5)),
        ("high_oversight_completions", close_abs(rnum("high_oversight_completions", "high_risk_monitored", "high_tier_monitored"), ref["high_oversight_completions"], 0.5)),
    ]
    rok = sum(1 for _, ok in roll_checks if ok)
    e += rok
    p += len(roll_checks)
    lines.append(f"    {rok:>3}/{len(roll_checks):<3} summary rollup metrics correct")

    # presence: every trial present in trial_risk_scores
    present = min(len(by_id), len(R))
    e += present
    p += len(R)
    lines.append(f"    {present:>3}/{len(R):<3} trial_risk_scores: every portfolio trial present")

    # capacity allocation consistency: no week over cap, total assigned matches reference
    capalloc = read_sheet(wb, "capacity_allocation") or []
    cap = ref["rb"]["weekly_cra_hours_cap"]
    total_assigned = 0
    weeks_ok = True
    for row in capalloc:
        h = num(get(row, "assigned_hours", "assigned", "hours"))
        if h is None:
            continue
        total_assigned += h
        if h > cap + 0.5:
            weeks_ok = False
    cap_checks = 0
    if capalloc and weeks_ok:
        cap_checks += 1
    if close_abs(total_assigned, ref["total_monitor_hours"], 1.0):
        cap_checks += 1
    e += cap_checks
    p += 2
    lines.append(f"    {cap_checks:>3}/{2:<3} CRA capacity allocation within cap and total hours consistent")

    return e, p, lines, by_id


def score_status_reconciliation(ref, wb):
    """The overall-vs-site status conflict log. This is a careful-reading deliverable:
    the divergences only surface when every record's per-site statuses are actually
    read and compared to the overall status. Each conflict trial earns up to 3 points
    -- identified, resolved to the governing overall status, and the divergent site
    values cited."""
    conflicts = read_sheet(wb, "conflict_log") or []
    cids = set(ref["conflict_ids"])
    gov = {c["nct_id"]: c["governing"] for c in ref["conflicts"]}
    seen = set()
    e = 0
    p = 3 * len(cids)
    ident = res = sites = 0
    for c in conflicts:
        nid = _nct(c)
        if not nid or nid in seen:
            continue
        seen.add(nid)
        if nid not in cids:
            continue
        ident += 1
        resolved = norm(get(c, "resolved_value", "governing_value", "resolved", "governing_status")).upper().replace(" ", "_")
        want = gov[nid].upper()
        if want in resolved or (resolved and resolved in want) or want.replace("_", "") in resolved.replace("_", ""):
            res += 1
        blob = " ".join(norm(v).lower() for v in c.values())
        # any genuine divergent site status word present proves they read the sites
        if any(w in blob for w in ("not_yet", "not yet", "completed", "withdrawn", "terminated",
                                   "active", "suspended", "enrolling", "available", "recruiting")):
            # require a 'site' framing so it is the divergence, not just the overall status
            if "site" in blob or "location" in blob or "facilit" in blob or get(c, "site_values", "site_statuses"):
                sites += 1
    e = ident + res + sites
    lines = [f"    {e:>3}/{p:<3} overall-vs-site status conflicts identified, resolved to governing value, and site values cited",
             f"        ({ident}/{len(cids)} identified, {res} resolved, {sites} with divergent site values)"]
    return e, p, lines


def score_source_register(sources, wb):
    rows = (read_sheet(wb, "source_register") or []) + (read_sheet(wb, "regulatory_basis") or [])
    e = 0
    p = 3 * len(sources)
    hit = 0
    for s in sources:
        sid = s["id"].lower().replace(" ", "")
        toks = s["fact_tokens"]
        doms = s["domains"]
        # First match the row LABELLED with this authority's source_id; only fall
        # back to a token scan when no row carries the id. This keeps each citation
        # tied to its own authority (loose fact tokens otherwise collide).
        row = None
        for r in rows:
            if norm(get(r, "source_id", "id", "source")).lower().replace(" ", "") == sid:
                row = r
                break
        if row is None:
            for r in rows:
                blob = " ".join(norm(v).lower() for v in r.values())
                if any(all(t in blob for t in grp) for grp in toks):
                    row = r
                    break
        if row is None:
            continue
        blob = " ".join(norm(v).lower() for v in row.values())
        fact_ok = any(all(t in blob for t in grp) for grp in toks)
        urls = URL_RE.findall(blob)
        dom_ok = any(any(d in u.lower() for d in doms) for u in urls)
        date_ok = bool(DATE_RE.search(blob))
        sc = (1 if fact_ok else 0) + (1 if dom_ok else 0) + (1 if date_ok else 0)
        e += sc
        if sc == 3:
            hit += 1
    lines = [f"    {e:>3}/{p:<3} authorities verified (fact token + official-domain URL + frozen date)",
             f"        ({hit}/{len(sources)} fully verified with all three)"]
    return e, p, lines


def score_trial_proofs(ref, report):
    active = ref["active_ids"]
    e = 0
    p = len(active)
    paras = re.split(r"\n\s*\n", report)
    for nid in active:
        ok = False
        for para in paras:
            if nid in para and len(para.strip()) >= JUSTIFY_MIN_CHARS and re.search(r"\d", para):
                ok = True
                break
        if not ok:
            for m in re.finditer(re.escape(nid), report):
                window = report[max(0, m.start() - 30): m.start() + JUSTIFY_MIN_CHARS + 60]
                if len(window) >= JUSTIFY_MIN_CHARS and re.search(r"\d", window.replace(nid, "")):
                    ok = True
                    break
        e += 1 if ok else 0
    lines = [f"    {e:>3}/{p:<3} in-scope trials with a quantified monitoring proof in the plan"]
    return e, p, lines


def score_audit_coverage(ref, report, wb):
    all_ids = ref["active_ids"] + ref["excluded_ids"]
    audit_rows = read_sheet(wb, "audit_log") or []
    by_id = {}
    for r in audit_rows:
        nid = _nct(r)
        if nid:
            by_id[nid] = r
    rlow = report.lower()
    e = 0
    p = len(all_ids)
    for nid in all_ids:
        ok = False
        row = by_id.get(nid)
        if row:
            blob = " ".join(norm(v).lower() for v in row.values())
            decision = norm(get(row, "decision", "disposition", "action", "outcome"))
            evidence = norm(get(row, "evidence_ref", "source_field", "source_cell", "citation", "evidence"))
            if (decision or evidence) and any(k in blob for k in DECISION_KEYWORDS):
                ok = True
        if not ok:
            i = rlow.find(nid.lower())
            if i >= 0:
                window = rlow[i: i + 320]
                if any(k in window for k in DECISION_KEYWORDS):
                    ok = True
        e += 1 if ok else 0
    lines = [f"    {e:>3}/{p:<3} trials carried to a decision in the audit log or plan"]
    return e, p, lines


def score_report_structure(ref, report):
    rlow = report.lower()
    e = 0
    p = 0
    lines = []

    sect_hits = sum(1 for grp in REPORT_SECTIONS if any(k in rlow for k in grp))
    e += sect_hits
    p += len(REPORT_SECTIONS)
    lines.append(f"    {sect_hits:>3}/{len(REPORT_SECTIONS):<3} required plan sections present")

    disc = sum(1 for d in KEY_DISCLOSURES if d in rlow)
    e += disc
    p += len(KEY_DISCLOSURES)
    lines.append(f"    {disc:>3}/{len(KEY_DISCLOSURES):<3} key triage concepts present")

    consistency = 0
    if str(ref["high_oversight_completions"]) in rlow or str(ref["window_infeasible_count"]) in rlow:
        consistency += 1
    if any(w in rlow for w in ("window", "capacity", "defer")) and any(w in rlow for w in ("infeasib", "oversubscrib", "exceed", "cadence", "defer")):
        consistency += 1
    e += consistency
    p += 2
    lines.append(f"    {consistency:>3}/{2:<3} plan numbers consistent with the recomputed schedule")

    no_ph = 1 if (rlow and not any(b in rlow for b in BANNED_TOKENS)) else 0
    e += no_ph
    p += 1
    lines.append(f"    {no_ph:>3}/{1:<3} no placeholder text")

    length_ok = 1 if len(report) >= REPORT_MIN_CHARS else round(len(report) / REPORT_MIN_CHARS, 2)
    length_ok = min(length_ok, 1)
    e += length_ok
    p += 1
    lines.append(f"    {round(length_ok,2):>3}/{1:<3} plan length (>= {REPORT_MIN_CHARS} chars; got {len(report)})")

    return e, p, lines


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--agent-output", default="/logs/agent/output.xlsx")
    ap.add_argument("--agent-dir", default="/logs/agent")
    ap.add_argument("--report", default="/logs/agent/monitoring_plan.md")
    ap.add_argument("--reward-out", default="/logs/verifier/reward.txt")
    args = ap.parse_args()

    def finish(score, lines):
        score = round(min(float(score), 1.0), 4)
        try:
            os.makedirs(os.path.dirname(args.reward_out), exist_ok=True)
            with open(args.reward_out, "w") as f:
                f.write(str(score))
        except OSError as ex:
            print(f"WARN: could not write reward: {ex}")
        try:
            with open(os.path.join(args.agent_dir, "judge_justification.txt"), "w") as f:
                f.write(f"Score: {score}\n\n" + "\n".join(lines))
        except OSError:
            pass
        print(f"Score: {score}")
        print("\n".join(lines))

    ref = build_reference()
    sources = load_sources()

    try:
        wb = openpyxl.load_workbook(args.agent_output, data_only=True, read_only=True)
    except Exception as ex:
        wb = None
        wb_err = str(ex)
    else:
        wb_err = None

    report = ""
    # Accept either the canonical name or a couple of obvious fallbacks.
    report_path = args.report
    if not os.path.exists(report_path):
        for alt in ("monitoring_plan.md", "scheduler_brief.md", "report.md"):
            cand = os.path.join(args.agent_dir, alt)
            if os.path.exists(cand):
                report_path = cand
                break
    if os.path.exists(report_path):
        try:
            with open(report_path, encoding="utf-8", errors="ignore") as f:
                report = f.read()
        except OSError:
            report = ""

    cats = {}
    detail = []

    if wb is not None:
        e, p, l, _ = score_numeric_backbone(ref, wb)
        cats["numeric_backbone"] = (e, p)
        detail.append(("numeric_backbone", l))
        e, p, l = score_status_reconciliation(ref, wb)
        cats["status_reconciliation"] = (e, p)
        detail.append(("status_reconciliation", l))
        e, p, l = score_source_register(sources, wb)
        cats["source_register"] = (e, p)
        detail.append(("source_register", l))
        e, p, l = score_audit_coverage(ref, report, wb)
        cats["audit_coverage"] = (e, p)
        detail.append(("audit_coverage", l))
    else:
        for k in ("numeric_backbone", "status_reconciliation", "source_register", "audit_coverage"):
            cats[k] = (0, 1)
        detail.append(("workbook", [f"    no usable workbook at {args.agent_output}: {wb_err}"]))

    e, p, l = score_trial_proofs(ref, report)
    cats["trial_proofs"] = (e, p)
    detail.append(("trial_proofs", l))
    e, p, l = score_report_structure(ref, report)
    cats["report_structure"] = (e, p)
    detail.append(("report_structure", l))

    score = 0.0
    cat_lines = []
    for cat, w in WEIGHTS.items():
        e, p = cats.get(cat, (0, 1))
        ratio = (e / p) if p else 0.0
        score += w * ratio
        cat_lines.append(f"  [{w:.2f}] {cat:18} {round(ratio,4):>6}  ({round(e,2)}/{p})")

    lines = [
        f"Clinical-trial risk-based monitoring weighted score = {round(score,4)}",
        f"reference: in_scope={len(ref['active_ids'])} excluded={len(ref['excluded_ids'])}"
        f" high={ref['high_tier_count']} medium={ref['medium_tier_count']} low={ref['low_tier_count']}"
        f" regulated={ref['regulated_count']} status_conflicts={ref['status_conflicts_count']}"
        f" window_infeasible={ref['window_infeasible_count']} capacity_deferred={ref['capacity_deferred_count']}"
        f" scheduled={ref['scheduled_count']} high_oversight_completions={ref['high_oversight_completions']}",
        f"workbook sheets present: {[s for s in REQUIRED_SHEETS if wb is not None and s in wb.sheetnames]}",
        "",
        "category weighted ratios:",
    ] + cat_lines + ["", "detail:"]
    for _cat, l in detail:
        lines.extend(l)
    finish(score, lines)


if __name__ == "__main__":
    main()

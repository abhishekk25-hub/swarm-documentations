#!/usr/bin/env python3
import argparse
import csv
import glob
import json
import os
import re
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
CORPUS = os.path.join(HERE, "corpus")
COHORT = os.path.join(HERE, "cohort.json")
API = "https://api.fireworks.ai/inference/v1/chat/completions"
MODEL = "accounts/fireworks/models/gpt-oss-120b"
W = {"answer": 0.50, "method": 0.20, "relation": 0.20, "topic": 0.10}
PROBLEMS_WEIGHT = 0.90
DELIVERABLES_WEIGHT = 0.10
MAXWORKERS = 4
TITLE_STOP = set("aime problem i ii the of and a an".split())

SYS = ("You grade a worker's solved answer key for United States AIME competition problems. You "
       "receive one problem statement, its official answer, and the values the worker recorded. Judge "
       "strictly from the statement and the official answer, then reply with only a JSON object.")


def norm(s):
    s = unicodedata.normalize("NFKC", str(s or ""))
    return re.sub(r"\s+", " ", s).strip().lower()


def title_tokens(s):
    return {w for w in re.findall(r"[a-z0-9]+", norm(s)) if w not in TITLE_STOP}


def col(row, *keys):
    for k in row:
        if any(t in k for t in keys):
            return k
    return None


def _rows_from(table):
    if not table:
        return []
    hdr = [norm(h) for h in table[0]]
    out = []
    for raw in table[1:]:
        if raw is None or all(c is None or norm(c) == "" for c in raw):
            continue
        out.append({hdr[i]: ("" if i >= len(raw) or raw[i] is None else str(raw[i]).strip())
                    for i in range(len(hdr))})
    return out


def _workbook(agent_dir):
    xp = os.path.join(agent_dir, "problems.xlsx")
    if os.path.isfile(xp):
        try:
            return openpyxl.load_workbook(xp, data_only=True, read_only=True)
        except Exception:
            return None
    return None


def _rows_from_records(agent_dir):
    rows = []
    for p in sorted(glob.glob(os.path.join(agent_dir, "records", "*.json"))):
        try:
            d = json.load(open(p, encoding="utf-8", errors="replace"))
        except Exception:
            continue
        if isinstance(d, dict):
            rows.append({norm(k): ("" if v is None else str(v).strip()) for k, v in d.items()})
    return rows


def read_table(agent_dir):
    wb = _workbook(agent_dir)
    if wb is not None:
        sheet = next((n for n in wb.sheetnames if "problem" in norm(n)), wb.sheetnames[0])
        rows = _rows_from(list(wb[sheet].iter_rows(values_only=True)))
        if rows:
            return rows
    cands = ([os.path.join(agent_dir, n) for n in ("problems.csv", "problems.xlsx.csv")]
             + sorted(glob.glob(os.path.join(agent_dir, "*problem*.csv"))))
    for p in cands:
        if os.path.isfile(p):
            try:
                with open(p, newline="", encoding="utf-8", errors="replace") as fh:
                    rows = _rows_from([r for r in csv.reader(fh)])
                if rows:
                    return rows
            except Exception:
                continue
    return _rows_from_records(agent_dir)


def read_summary(agent_dir):
    wb = _workbook(agent_dir)
    if wb is not None:
        sn = next((n for n in wb.sheetnames if "summ" in norm(n)), None)
        if sn:
            rows = list(wb[sn].iter_rows(values_only=True))
            txt = "\n".join(" | ".join("" if c is None else str(c) for c in r) for r in rows if r)
            if txt.strip():
                return txt[:2000]
    return "(no summary sheet provided)"


def corpus_text(rid):
    p = os.path.join(CORPUS, rid + ".txt")
    try:
        return open(p, encoding="utf-8", errors="replace").read()
    except OSError:
        return ""


def read_brief(agent_dir):
    for nm in ("brief.md", "memo.md", "report.md"):
        p = os.path.join(agent_dir, nm)
        if os.path.isfile(p):
            return open(p, encoding="utf-8", errors="ignore").read()
    return ""


def call_model(user, key):
    payload = {"model": MODEL, "temperature": 0, "max_tokens": 4000,
               "messages": [{"role": "system", "content": SYS}, {"role": "user", "content": user}]}
    last = ""
    for attempt in range(5):
        try:
            r = requests.post(API, headers={"Authorization": "Bearer " + key,
                              "Content-Type": "application/json"}, json=payload, timeout=180)
            if r.status_code == 200:
                return r.json()["choices"][0]["message"]["content"]
            last = "http %s" % r.status_code
            if r.status_code not in (429, 500, 502, 503, 504):
                break
        except Exception as e:
            last = str(e)[:80]
        time.sleep(2 * (attempt + 1))
    raise RuntimeError(last or "call failed")


def parse_obj(content):
    obj = None
    for cand in re.findall(r"\{[^{}]*\}", content, re.S):
        try:
            obj = json.loads(cand)
        except ValueError:
            continue
    if obj is None:
        raise ValueError("no json")
    return obj


def judge_problem(text, official, row, key):
    ac = col(row, "final_answer", "answer")
    mc = col(row, "method")
    rc = col(row, "relation", "key", "result")
    tc = col(row, "topic", "area", "subject")
    user = (
        "PROBLEM STATEMENT:\n" + text
        + "\n\nOFFICIAL ANSWER (reference, a whole number 0-999): " + str(official) + "\n\n"
        "RECORDED VALUES:\n"
        + "final_answer: " + str(row.get(ac, "") or "(none recorded)") + "\n"
        + "principal_method: " + str(row.get(mc, "") or "(none recorded)") + "\n"
        + "key_relation: " + str(row.get(rc, "") or "(none recorded)") + "\n"
        + "topic: " + str(row.get(tc, "") or "(none recorded)") + "\n\n"
        "Decide the following, judging only from the statement and the official answer.\n"
        "answer_correct: true if the recorded final_answer is the same whole number as the official "
        "answer for this problem, comparing the two as integers in 0..999 and ignoring leading zeros, "
        "surrounding words, or formatting; false if it differs, is missing, or is not a number.\n"
        "method_correct: true if principal_method names an approach that genuinely yields the official "
        "answer for this problem (a correct method, not necessarily the only one); false if it is wrong, "
        "empty, or generic boilerplate that does not fit the problem.\n"
        "relation_correct: true if key_relation states an equation, identity, or count that is "
        "mathematically correct for this problem and is genuinely central to reaching the official "
        "answer; false if it is wrong, irrelevant, empty, or a mere restatement of the question.\n"
        "true_topic: the best single fit among algebra, number_theory, geometry, combinatorics.\n"
        'Reply with only this JSON: {"answer_correct": <bool>, "method_correct": <bool>, '
        '"relation_correct": <bool>, "true_topic": "algebra|number_theory|geometry|combinatorics"}.')
    o = parse_obj(call_model(user, key))
    return {
        "answer_correct": bool(o.get("answer_correct")),
        "method_correct": bool(o.get("method_correct")),
        "relation_correct": bool(o.get("relation_correct")),
        "true_topic": norm(o.get("true_topic")) or "algebra",
        "recorded_topic": norm(row.get(tc, "")),
    }


def judge_deliverables(brief, summary, key):
    if len(brief.strip()) < 200 and summary == "(no summary sheet provided)":
        return 0.0
    user = (
        "A worker produced a short brief and a summary sheet for a set of AIME competition problems they "
        "solved. Judge them on their own terms.\n\nBRIEF:\n" + brief[:7000]
        + "\n\nSUMMARY SHEET ROWS:\n" + summary + "\n\n"
        "grounded: true if the brief discusses specific problems, methods, or topics rather than generic "
        "filler.\n"
        "breakdown: true if the brief describes how the set breaks down across topics or difficulty.\n"
        "summary_ok: true if the summary sheet gives a sensible rollup such as the total count and counts "
        "per topic.\n"
        'Reply with only this JSON: {"grounded": <bool>, "breakdown": <bool>, "summary_ok": <bool>}.')
    o = parse_obj(call_model(user, key))
    return 0.45 * bool(o.get("grounded")) + 0.30 * bool(o.get("breakdown")) + 0.25 * bool(o.get("summary_ok"))


def balanced(verds, true_key, ok_fn):
    counts = {}
    for v in verds:
        counts[v[true_key]] = counts.get(v[true_key], 0) + 1
    k = max(1, len(counts))
    return sum(1.0 / (k * counts[v[true_key]]) for v in verds if ok_fn(v))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--agent-dir", default="/logs/agent")
    ap.add_argument("--reward-out", default="/logs/verifier/reward.txt")
    a = ap.parse_args()
    lines = []

    def finish(score):
        score = round(max(0.0, min(1.0, score)), 4)
        os.makedirs(os.path.dirname(a.reward_out) or ".", exist_ok=True)
        open(a.reward_out, "w").write(str(score))
        body = "\n".join(["Score: %s" % score, ""] + lines)
        for p in ("/logs/verifier/verify_justification.txt", os.path.join(a.agent_dir, "verify_justification.txt")):
            try:
                os.makedirs(os.path.dirname(p), exist_ok=True)
                open(p, "w").write(body)
            except OSError:
                pass
        print(body)

    key = os.environ.get("FIREWORKS_API_KEY") or os.environ.get("OPENAI_API_KEY")
    if not key:
        lines.append("No judge API key in environment; failing closed.")
        finish(0.0)
        return
    try:
        cohort = json.load(open(COHORT))
    except (OSError, ValueError):
        lines.append("No cohort.json.")
        finish(0.0)
        return
    rows = read_table(a.agent_dir)
    if not rows:
        lines.append("No usable problems table (.xlsx, .csv, or records/) in agent dir.")
        finish(0.0)
        return

    id_col = col(rows[0], "problem_id", "id")
    title_col = col(rows[0], "title", "problem", "name")
    by_id, title_rows = {}, []
    for r in rows:
        rid = norm(r.get(id_col, "")) if id_col else ""
        if rid:
            by_id.setdefault(rid, r)
        if title_col:
            title_rows.append((title_tokens(r.get(title_col, "")), r))

    def match(entry):
        if norm(entry.get("id", "")) in by_id:
            return by_id[norm(entry["id"])]
        gt = title_tokens(entry.get("title", ""))
        best, bs, bo = None, 0, 0.0
        for at, r in title_rows:
            if not at or not gt:
                continue
            shared = len(at & gt)
            ov = shared / min(len(at), len(gt))
            if (shared, ov) > (bs, bo):
                bs, bo, best = shared, ov, r
        return best if (bs >= 2 or (bs >= 1 and bo >= 0.5)) else None

    def work(entry):
        text = corpus_text(entry["id"])
        if not text:
            return None
        row = match(entry) or {}
        try:
            v = judge_problem(text, entry["answer"], row, key)
        except Exception:
            return "FAIL"
        return v

    out = list(ThreadPoolExecutor(max_workers=MAXWORKERS).map(work, cohort))
    verds = [v for v in out if isinstance(v, dict)]
    failed = sum(1 for v in out if v == "FAIL")
    if failed > max(2, len(cohort) // 10):
        lines.append("Judge unavailable for %d/%d problems; failing closed." % (failed, len(cohort)))
        finish(0.0)
        return

    n = len(cohort)
    answer = sum(1 for v in verds if v["answer_correct"]) / n
    method = sum(1 for v in verds if v["method_correct"]) / n
    relation = sum(1 for v in verds if v["relation_correct"]) / n
    topic = balanced(verds, "true_topic", lambda v: v["recorded_topic"] == v["true_topic"])
    pr = W["answer"] * answer + W["method"] * method + W["relation"] * relation + W["topic"] * topic

    deliv_error = ""
    try:
        deliv = judge_deliverables(read_brief(a.agent_dir), read_summary(a.agent_dir), key)
    except Exception as e:
        deliv = 0.0
        deliv_error = str(e)[:160]

    reward = PROBLEMS_WEIGHT * pr + DELIVERABLES_WEIGHT * deliv
    lines += [
        "problems %.4f x %.2f + deliverables %.4f x %.2f" % (pr, PROBLEMS_WEIGHT, deliv, DELIVERABLES_WEIGHT),
        "answer %.3f | method %.3f | relation %.3f | balanced topic %.3f" % (answer, method, relation, topic),
        "cohort: %d; judged: %d; judge failures: %d" % (n, len(verds), failed),
    ]
    if deliv_error:
        lines.append("ERROR deliverables judge call failed (infrastructure, scored 0.0): " + deliv_error)
    finish(reward)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
Verifier for the Cross-Utility WMP Reliability Briefing task.

Grades ONLY the agent's generated deliverables under /logs/agent/:
  - cross_utility_findings.csv  (the grounded findings ledger; per-utility
    findings_<utility>.csv shards are merged in if present)
  - reliability_briefing.docx   (the written briefing, the centerpiece)
  - chart_*.png                 (ledger-derived charts)
  - briefing_deck.pptx          (the meeting deck)

It NEVER inspects the agent's trajectory, tool/fetch logs, session JSON, or how
the result was produced — the agent may use any method/tool. Only the output
files are scored.

Scoring has two parts, both read ONLY from the deliverables:
  - Checks 1-21: an LLM judge marks structural / quality / consistency criteria. The
    judge is run NUM_JUDGE_RUNS times independently and each check is decided by
    majority vote, so no single judge call decides the majority of the reward.
  - Checks 22-30: nine deterministic DEPTH-OF-COVERAGE checks (graduated threshold
    bands) computed in Python directly from the findings ledger. These reward the
    sustained PER-DECISION reading the instruction asks for (several findings per
    Decision across all 18 documents, spread across utilities and topics, each backed
    by a substantial verbatim quote) - the work a single context cannot keep up across
    18 heavy documents but a swarm of per-Decision readers can. This is the genuine
    multi-vs-single separation, not a trajectory check.

Anti-padding guard: every depth count is taken over SUBSTANTIVE findings only (see
is_substantive_finding). A row counts as a genuine finding only if regulator_finding
is a real paraphrase SENTENCE asserting what the regulator did/required - not a section
heading, not a directive title or label/ID (e.g. "PG&E-23B-01"), not a table-of-contents
line with dot leaders / page numbers, and not text copied verbatim out of the document.
This closes two scrapes: (1) flooding the ledger with TOC entries or quote slices, and
(2) the "title in finding, body in quote" split, where the agent lifts a Decision's
numbered directive title into the finding and its body into the quote - two adjacent
verbatim fragments that fool a finding-inside-quote test but are not a paraphrase plus
evidence. Such rows are dropped before counting and are also flagged to the LLM judge
(raw vs. substantive counts plus a [GENUINE]/[PADDING] tag on every sampled row), so
they cannot move either half of the score.

Scoring is simple and additive: one point per check, every check weighted equally.
  reward = (checks passed) / (total checks)
No weighting, no tiers, no caps. The gap comes from the task scale and coordination
requirement, not from how points are awarded.

Fail-closed: this is an llm-judge task, so the deterministic depth checks (22-30) never
produce a non-zero reward on their own. If every LLM judge call fails, the verifier
writes an explicit error state ({"reward": 0.0, "error": ...}) and the depth checks are
not counted - an infrastructure failure is never scored as a genuine zero-reward run.
"""

import csv
import io
import json
import os
import re
import traceback
from collections import Counter
from pathlib import Path

AGENT_DIR = Path("/logs/agent")
REWARD_PATH = Path("/logs/verifier/reward.json")
REWARD_TXT = Path("/logs/verifier/reward.txt")
LLM_CHECKS = 21
DEPTH_CHECKS = 9
NUM_CHECKS = 30
# The LLM judge is run this many times independently and majority-voted, so no
# single judge call decides the majority of the reward.
NUM_JUDGE_RUNS = 3

# Every check is worth one point. The reward is the simple fraction of checks passed
# (LLM quality checks 1-21 and deterministic depth-of-coverage checks 22-30 are all
# equal). The depth checks are not blocking gates: they run independently of and in
# addition to the LLM judge, so a depth miss never stops the LLM checks from scoring.

LEDGER_COLUMNS = [
    "finding_id", "utility", "doc_id", "topic", "regulator_finding", "quote", "source_url",
]
UTILITIES = ["PG&E", "SCE", "SDG&E", "BVES", "Liberty", "PacifiCorp", "HWT", "TBC", "LS Power"]
CHART_FILES = [
    "chart_findings_by_topic.png",
    "chart_findings_by_utility.png",
    "chart_directive_themes.png",
]

# The 18 WMP Decisions (a draft + a final per utility). Maps the e-filing fileid
# (used as doc_id) to its utility and version. Used only to measure how much of
# the document set the ledger actually draws on.
DOC_INFO = {
    "57273": ("PG&E", "draft"),       "57629": ("PG&E", "final"),
    "57232": ("SCE", "draft"),        "57548": ("SCE", "final"),
    "57229": ("SDG&E", "draft"),      "57541": ("SDG&E", "final"),
    "57256": ("BVES", "draft"),       "57526": ("BVES", "final"),
    "57863": ("Liberty", "draft"),    "58230": ("Liberty", "final"),
    "57969": ("PacifiCorp", "draft"), "58352": ("PacifiCorp", "final"),
    "57492": ("HWT", "draft"),        "57706": ("HWT", "final"),
    "57493": ("TBC", "draft"),        "57707": ("TBC", "final"),
    "57730": ("LS Power", "draft"),   "57985": ("LS Power", "final"),
}


def _norm_topic(t):
    return re.sub(r"[^a-z0-9]", "", (t or "").lower())


def _norm_text(s):
    return re.sub(r"\s+", " ", (s or "").strip().lower())


# A genuine paraphrase of an Energy Safety finding states what the regulator did/required,
# so it almost always carries one of these action verbs. A scraped section TITLE
# ("Cross-Utility Collaboration on Risk Model Development") is a noun phrase with none.
_FINDING_VERB_RE = re.compile(
    r"\b("
    r"require\w*|must|shall|should|direct\w*|order\w*|find|finds|found|determine\w*|"
    r"note|notes|noted|identif\w+|approv\w+|acknowledg\w+|accept\w*|criticiz\w+|"
    r"conclud\w+|expect\w*|need\w*|provide\w*|submit\w*|report\w*|demonstrate\w*|"
    r"ensure\w*|address\w*|evaluate\w*|recommend\w*|flag\w*|fail\w*|lack\w*|"
    r"improve\w*|reduce\w*|increase\w*|continue\w*|allow\w*|deny|denied|reject\w*|"
    r"is required|are required|was required|were required|did not|does not|has not|have not"
    r")\b",
    re.IGNORECASE,
)

# Start-anchored directive label/ID the scrape prepends, e.g. "PG&E - PG&E-23B-01:" or
# "SCE-22-09 ...". Anchored to the start so a code mentioned mid-sentence in a real
# paraphrase does not trip it.
_CODE_PREFIX_RE = re.compile(
    r"^\s*[A-Za-z][\w&.\-]{0,12}(\s*[-\u2013:]\s*[A-Za-z&.]{1,12})?[-\u2013]\d{1,4}[A-Za-z]?\b",
    re.IGNORECASE,
)


def _title_case_ratio(s):
    """Fraction of words that begin with a capital letter. Section titles are mostly
    Title-Cased ('Deployment of New Technologies'); real sentences are not."""
    words = re.findall(r"[A-Za-z][A-Za-z'\u2019&.\-]*", s)
    if len(words) < 3:
        return 1.0
    cap = sum(1 for w in words if w[:1].isupper())
    return cap / len(words)


def is_substantive_finding(r):
    """A row counts as a GENUINE finding only if its regulator_finding is a real
    paraphrase SENTENCE in the analyst's own words - not scraped boilerplate. The
    instruction is explicit that the finding must be the regulator's point in the
    agent's own words and DIFFERENT text from the quote, so these are content
    requirements, not formatting nit-picks.

    This guard closes two distinct scrapes seen in practice:
      1. Table-of-contents / heading padding (dot-leaders, page numbers).
      2. The "title in finding, body in quote" split: the agent lifts a Decision's
         numbered directive TITLE (e.g. "PG&E-23B-01: Cross-Utility Collaboration on
         Risk Model Development") into the finding and the directive BODY into the
         quote. Because the two are different (adjacent) verbatim fragments, an
         is-the-finding-inside-the-quote test alone is fooled - so we additionally
         require the finding to read like a genuine sentence.

    A row is substantive only when ALL hold:
      - regulator_finding is present, >= 25 chars and >= 6 words;
      - it has no dot-leader run (".... 12");
      - it is not identical to, nor a verbatim slice of, the quote;
      - it does not begin with a directive label/ID (e.g. "PG&E-23B-01");
      - it contains a regulator action verb (it asserts what the regulator did);
      - it is not a mostly Title-Cased heading (capitalized-word ratio < 0.6).

    It ALSO requires a genuine supporting QUOTE. The instruction is explicit that a
    finding without a real verbatim quote from the decision is useless and should be
    left out, so a finding whose "quote" is just a table-of-contents heading (dot-leaders
    + page number, e.g. "PG&E-23B-01. Cross-Utility Collaboration ...... 12") or a bare
    Title-Cased directive title is NOT substantive - that is the latest scrape, where the
    paraphrase is a mechanical "[Utility] must address [ID] regarding [topic]" template
    pointing at a TOC line rather than at real document prose. The quote must therefore be
    a real passage:
      - present and at least 40 characters of actual text;
      - no dot-leader run (TOC artifact);
      - not a mostly Title-Cased heading line.
    """
    f = (r.get("regulator_finding") or "").strip()
    q = (r.get("quote") or "").strip()
    if len(f) < 25:
        return False
    if len(re.findall(r"[A-Za-z]+", f)) < 6:
        return False
    if re.search(r"\.\s*\.\s*\.\s*\.", f):  # dot-leader / TOC artifact
        return False
    nf, nq = _norm_text(f), _norm_text(q)
    if not nf or nf == nq or nf in nq:
        return False
    if _CODE_PREFIX_RE.search(f):           # starts with a directive label/ID
        return False
    if not _FINDING_VERB_RE.search(f):      # no action verb -> a noun-phrase title
        return False
    if _title_case_ratio(f) >= 0.6:         # mostly Title-Cased heading
        return False
    # Quote must be a genuine verbatim passage, not a TOC heading or directive title.
    if len(q) < 40:
        return False
    if re.search(r"\.\s*\.\s*\.\s*\.", q):  # dot-leader / TOC line in the quote
        return False
    if _title_case_ratio(q) >= 0.6:         # quote is a Title-Cased heading, not prose
        return False
    return True


def compute_depth_metrics(rows):
    """Objective depth-of-coverage metrics, computed from the output ledger only and
    counting ONLY substantive (genuine) findings - see is_substantive_finding.

    The separating signal is PER-DECISION depth: the instruction asks for several
    distinct findings per Decision across the whole 18-document set ("don't let it thin
    out after the first few"). A complete-but-shallow pass can name every utility and
    cite every Decision once, but it cannot put real depth behind each one. Crucially,
    a run cannot fake depth by flooding the ledger with scraped table-of-contents lines
    or quote slices either: those rows fail is_substantive_finding and are dropped
    before any depth count, so padding does not move the score. These metrics measure
    GENUINE depth - how thoroughly each Decision was actually mined - not raw row count."""
    raw_total = len(rows)
    sub = [r for r in rows if is_substantive_finding(r)]
    total = len(sub)
    doc_ids = [re.sub(r"\D", "", (r.get("doc_id") or "")) for r in sub]
    known_doc_ids = set(d for d in doc_ids if d in DOC_INFO)

    util_counts = Counter((r.get("utility") or "").strip() for r in sub if (r.get("utility") or "").strip())
    topic_counts = Counter(_norm_topic(r.get("topic")) for r in sub if _norm_topic(r.get("topic")))

    # findings per individual Decision (per known doc_id) - the per-document depth signal
    doc_counts = Counter(d for d in doc_ids if d in DOC_INFO)

    long_quotes = sum(1 for r in sub if len((r.get("quote") or "").strip()) >= 80)

    return {
        "raw_rows": raw_total,
        "substantive_findings": total,
        "total_findings": total,
        "distinct_known_doc_ids": len(known_doc_ids),
        "decisions_with_ge3_findings": sum(1 for c in doc_counts.values() if c >= 3),
        "decisions_with_ge5_findings": sum(1 for c in doc_counts.values() if c >= 5),
        "utilities_with_ge10_findings": sum(1 for c in util_counts.values() if c >= 10),
        "topics_with_ge8_findings": sum(1 for c in topic_counts.values() if c >= 8),
        "findings_with_quote_ge80_chars": long_quotes,
    }


def depth_checks(rows):
    """Checks 22-30: deterministic depth-of-coverage, graduated threshold bands, counting
    ONLY substantive findings (see is_substantive_finding). Each is one independent point
    (no weighting). They reward the sustained per-Decision reading that the instruction
    asks for across all 18 documents - the work that does not fit a single context but
    spreads naturally across a swarm of per-Decision readers. Because padding rows are
    filtered out before counting, a run cannot pass these by scraping headings or copying
    quote slices; it has to actually produce many genuine findings. A shallow run still
    keeps full credit on the quality checks (1-21); these only measure depth."""
    m = compute_depth_metrics(rows)
    specs = [
        (22, m["decisions_with_ge3_findings"] >= 10,
         f"{m['decisions_with_ge3_findings']} of the 18 Decisions have >= 3 substantive findings each "
         f"(need >= 10; several genuine findings per Decision, not one or two)."),
        (23, m["decisions_with_ge3_findings"] >= 14,
         f"{m['decisions_with_ge3_findings']} of the 18 Decisions have >= 3 substantive findings each "
         f"(need >= 14 for near-full-set depth)."),
        (24, m["decisions_with_ge5_findings"] >= 6,
         f"{m['decisions_with_ge5_findings']} of the 18 Decisions have >= 5 substantive findings each "
         f"(need >= 6; real depth on individual Decisions)."),
        (25, m["decisions_with_ge5_findings"] >= 10,
         f"{m['decisions_with_ge5_findings']} of the 18 Decisions have >= 5 substantive findings each "
         f"(need >= 10 for sustained depth across the set)."),
        (26, m["total_findings"] >= 90,
         f"Ledger has {m['total_findings']} substantive findings "
         f"(need >= 90; raw rows = {m['raw_rows']})."),
        (27, m["total_findings"] >= 150,
         f"Ledger has {m['total_findings']} substantive findings "
         f"(need >= 150 for full-fleet depth; raw rows = {m['raw_rows']})."),
        (28, m["utilities_with_ge10_findings"] >= 5,
         f"{m['utilities_with_ge10_findings']} utilities have >= 10 substantive findings each "
         f"(need >= 5; depth must spread beyond the largest two utilities)."),
        (29, m["topics_with_ge8_findings"] >= 5,
         f"{m['topics_with_ge8_findings']} topics have >= 8 substantive findings each "
         f"(need >= 5 of the 8 topics for taxonomy-wide depth)."),
        (30, m["findings_with_quote_ge80_chars"] >= 100,
         f"{m['findings_with_quote_ge80_chars']} substantive findings carry a substantial verbatim "
         f"quote (>= 80 chars) (need >= 100)."),
    ]
    return [{"id": cid, "pass": ok, "reason": reason} for cid, ok, reason in specs], m


def _read_text(path):
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""


# The ledger schema is described in prose in instruction.md, so an agent may pick
# reasonable-but-different header names (e.g. document_id, finding, url, id). We grade
# the CONTENT the agent produced, not the exact header spelling, so equivalent columns
# are mapped onto the canonical schema below. This prevents a purely cosmetic naming
# choice from zeroing out genuinely grounded findings.
_CANON_ALIASES = {
    "finding_id":        {"id", "fid", "rowid", "rowno", "no", "findingno", "index", "idx"},
    "utility":           {"util", "corporation", "company", "electricalcorporation", "electriccorporation"},
    "doc_id":            {"docid", "documentid", "decisionid", "fileid", "efilingid", "efilingfileid", "documentnumber"},
    "topic":             {"category", "theme", "topictag", "topicarea"},
    "regulator_finding": {"regulatoryfinding", "finding", "regulatorpoint", "point", "paraphrase", "summary",
                          "findingtext", "regulatordirective", "directive", "findingsummary"},
    "quote":             {"verbatim", "verbatimquote", "evidence", "evidencequote", "excerpt", "passage", "sourcequote"},
    "source_url":        {"url", "link", "sourcelink", "source", "efilingurl", "decisionurl", "documenturl"},
}


def _canon_key(name):
    return re.sub(r"[^a-z0-9]", "", (name or "").lower())


_ALIAS_TO_CANON = {}
for _canon, _aliases in _CANON_ALIASES.items():
    _ALIAS_TO_CANON[_canon_key(_canon)] = _canon
    for _a in _aliases:
        _ALIAS_TO_CANON[_canon_key(_a)] = _canon


def _normalize_row(r):
    """Map a row's columns onto the canonical ledger schema regardless of the exact
    header spelling. An explicit canonical column always wins; an alias only fills a
    canonical field that is otherwise empty."""
    out = dict(r)
    for raw_key, val in r.items():
        canon = _ALIAS_TO_CANON.get(_canon_key(raw_key))
        if not canon:
            continue
        if (out.get(canon) or "").strip():
            continue
        if (val or "").strip():
            out[canon] = val
    return out


def _parse_csv_rows(text):
    if not text.strip():
        return []
    try:
        reader = csv.DictReader(io.StringIO(text))
        return [_normalize_row(dict(r)) for r in reader]
    except Exception:
        return []


def load_ledger():
    """Load the findings ledger from cross_utility_findings.csv, merging any
    per-utility findings_<utility>.csv shards. De-duplicates on (doc_id, quote).
    Only output files are read — no trajectory inspection."""
    rows = []
    rows.extend(_parse_csv_rows(_read_text(AGENT_DIR / "cross_utility_findings.csv")))
    try:
        shards = sorted(
            fn for fn in os.listdir(AGENT_DIR)
            if fn.lower().startswith("findings_") and fn.lower().endswith(".csv")
        )
    except OSError:
        shards = []
    for fn in shards:
        rows.extend(_parse_csv_rows(_read_text(AGENT_DIR / fn)))

    seen, merged = set(), []
    for r in rows:
        key = ((r.get("doc_id") or "").strip(), (r.get("quote") or "").strip()[:120])
        if key == ("", ""):
            continue
        if key in seen:
            continue
        seen.add(key)
        merged.append(r)
    return merged


def summarize_ledger(rows):
    if not rows:
        return "[no findings ledger found: cross_utility_findings.csv and findings_*.csv are missing or empty]"

    header = list(rows[0].keys())
    sub_flags = [is_substantive_finding(r) for r in rows]
    n_sub = sum(sub_flags)
    parts = [
        f"Total findings rows (after merging shards + de-dup): {len(rows)}",
        f"Of those, SUBSTANTIVE (genuine paraphrase findings, not TOC lines / headings / "
        f"text copied out of the quote): {n_sub}  "
        f"[{len(rows) - n_sub} rows look like non-substantive padding]",
        "Columns present: " + ", ".join(header),
    ]

    from collections import Counter
    sub_rows = [r for r, ok in zip(rows, sub_flags) if ok]
    util_counts = Counter((r.get("utility") or "").strip() for r in sub_rows)
    topic_counts = Counter((r.get("topic") or "").strip() for r in sub_rows)
    doc_counts = Counter((r.get("doc_id") or "").strip() for r in sub_rows)
    parts.append("SUBSTANTIVE findings per utility: " + json.dumps(dict(util_counts), ensure_ascii=False))
    parts.append("SUBSTANTIVE findings per topic: " + json.dumps(dict(topic_counts), ensure_ascii=False))
    parts.append(f"Distinct doc_ids cited by substantive findings: {len([d for d in doc_counts if d])}")

    n = len(rows)
    if n <= 30:
        sample_idx = list(range(n))
    else:
        sample_idx = list(range(0, 12)) + list(range(n // 2 - 3, n // 2 + 3)) + list(range(n - 12, n))
        sample_idx = sorted(set(i for i in sample_idx if 0 <= i < n))
    parts.append(
        f"\nSample of {len(sample_idx)} rows. Each is tagged [GENUINE] or [PADDING] by the "
        f"same substantive-finding rule (PADDING = the finding is a heading, a TOC line with "
        f"dot leaders, or text copied out of the quote):"
    )
    parts.append("(utility | doc_id | topic | regulator_finding | quote | source_url)")
    for i in sample_idx:
        r = rows[i]
        tag = "GENUINE" if sub_flags[i] else "PADDING"
        parts.append(
            f"  [{i+1}][{tag}] {r.get('utility','')} | {r.get('doc_id','')} | {r.get('topic','')} | "
            f"{(r.get('regulator_finding','') or '')[:260]} | "
            f"QUOTE: {(r.get('quote','') or '')[:300]} | {r.get('source_url','')}"
        )
    return "\n".join(parts)


def load_docx(path):
    if not path.exists():
        return "[reliability_briefing.docx not found]", 0
    try:
        import docx
        d = docx.Document(str(path))
    except Exception as exc:
        return f"[could not open reliability_briefing.docx: {exc}]", 0
    paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    text = "\n".join(paras)
    words = len(re.findall(r"\b\w+\b", text))
    return text[:48000], words


def load_pptx(path):
    if not path.exists():
        return "[briefing_deck.pptx not found]", 0
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
    except Exception as exc:
        return f"[could not open briefing_deck.pptx: {exc}]", 0
    slides_text = []
    for idx, slide in enumerate(prs.slides, 1):
        chunks = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text.strip())
        slides_text.append(f"--- Slide {idx} ---\n" + "\n".join(chunks))
    n = len(slides_text)
    return ("\n\n".join(slides_text))[:24000], n


def list_charts():
    present = []
    try:
        files = os.listdir(AGENT_DIR)
    except OSError:
        files = []
    for cf in CHART_FILES:
        if cf in files:
            present.append(cf)
    extra_pngs = [f for f in files if f.lower().endswith(".png") and f not in CHART_FILES]
    return present, extra_pngs


EVALUATION_PROMPT = """You are a strict but fair evaluator for a cross-utility grid-reliability briefing task.

An agent was asked to read across California Energy Safety's 2025 Wildfire
Mitigation Plan (WMP) Decisions for the state's electrical corporations — PG&E,
SCE, SDG&E, BVES, Liberty, PacifiCorp, HWT, TBC, and LS Power (each with a draft
and a final Decision) — and produce, under /logs/agent/, four deliverables that
all agree with each other:
  - cross_utility_findings.csv  (the grounded findings ledger; one row per
    finding with columns finding_id, utility, doc_id, topic, regulator_finding,
    quote, source_url)
  - reliability_briefing.docx   (the written briefing — the centerpiece)
  - chart_*.png                 (ledger-derived charts)
  - briefing_deck.pptx          (the meeting deck)

You are shown ONLY the contents of those output files (a summary of the ledger,
the briefing text, the deck text, and the list of chart files present). Judge
only what is in the deliverables — do not assume anything about how the agent
produced them. Each finding's `quote` should read like a genuine verbatim
passage lifted from a dense regulatory decision; you cannot re-open the source
documents, so judge the quotes on plausibility, substance, and distinctness
(real, varied regulatory sentences vs. short stubs or the same template repeated
with a utility name swapped in).

IMPORTANT - reject scraped/padded findings. A row only counts as a finding if its
regulator_finding is the regulator's point written as a REAL PARAPHRASE SENTENCE in
the analyst's own words. The following do NOT count as findings, even if the text is
different from the quote:
  - section headings or directive TITLES (e.g. "Deployment of New Technologies",
    "Cross-Utility Collaboration on Risk Model Development");
  - directive labels or IDs (e.g. "PG&E-23B-01", "SCE-22-09"), or a finding that
    begins with such a label;
  - table-of-contents lines (often with "....." dot leaders or a trailing page number);
  - text copied verbatim out of the document, including the case where the agent puts
    a directive's TITLE in regulator_finding and the directive's BODY in the quote -
    these are two adjacent verbatim fragments, NOT a finding plus its evidence.
A genuine regulator_finding reads like a sentence asserting what the regulator did or
required ("Energy Safety required SCE to ...", "The regulator found that PG&E ..."),
not a noun-phrase title. The ledger summary reports how many rows are SUBSTANTIVE vs
padding and tags every sampled row [GENUINE] or [PADDING]; base the ledger checks on
the SUBSTANTIVE findings, never on the raw row count, and independently apply the rule
above to the sampled rows you can see.

Score the output on these 21 independent boolean checks. Mark each PASS or FAIL.

FINDINGS LEDGER (cross_utility_findings.csv):
 1. A findings ledger exists with data rows and includes the columns finding_id, utility, doc_id, topic, regulator_finding, quote, and source_url (column names may vary slightly but all seven concepts must be present).
 2. The ledger has at least 30 SUBSTANTIVE finding rows (use the reported substantive count; padding/heading/title rows do not count toward this).
 3. At least 7 of the 9 utilities (PG&E, SCE, SDG&E, BVES, Liberty, PacifiCorp, HWT, TBC, LS Power) appear in the utility column of substantive findings.
 4. At least 5 distinct topics appear among substantive findings (from: grid hardening, asset inspection, equipment maintenance, vegetation management, situational awareness, emergency preparedness, PSPS, data / risk modeling).
 5. The doc_id column is populated for almost all rows with real-looking decision document ids (e.g. numeric e-filing ids like 57629, 57548), not blanks or invented labels.
 6. For at least 80% of the sampled rows, the quote field is a substantial passage that reads like a coherent sentence beginning at a sentence boundary (roughly 60+ chars) — FAIL if many quotes begin mid-sentence with a dangling title fragment (e.g. "Model Development In its Decision ...", "Transparency In its Decision ..."), which is the signature of a directive split between the finding and the quote.
 7. The vast majority of sampled rows are tagged [GENUINE], AND on your own reading regulator_finding is a real paraphrase sentence asserting what the regulator did/required — NOT a section title, a directive label/ID, or text lifted verbatim from the document. FAIL if a large share of the sample is tagged [PADDING], reads as a Title-Cased heading, or begins with a directive code like "PG&E-23B-01", even when it differs from the quote.
 8. The findings across the sampled rows are genuinely distinct in substance — not the same template repeated with a utility name/id swapped, and not a mechanical enumeration of every numbered directive heading in document order.
 9. For most sampled rows, source_url points to an Energy Safety / e-filing source (e.g. contains 'energysafety' or 'efiling') rather than a blank or an unrelated site.
10. Coverage is not lopsided: at least 5 utilities have 2 or more substantive findings each, and no single utility accounts for more than ~60% of all substantive findings.

RELIABILITY BRIEFING (reliability_briefing.docx):
11. The briefing exists and is a substantial written document (at least ~1500 words).
12. It opens with a portfolio overview describing what this set of WMP Decisions covers across the fleet.
13. It has a per-utility section that names specific utilities (at least 7 of the 9) and, for each, explains in the analyst's own words what that utility was directed to fix or where it fell short — genuine analysis, NOT a pasted list of that utility's directive titles/IDs.
14. The per-utility paragraphs are genuinely distinct analytical prose — each says something specific to that utility in sentences, not the same sentence with the name swapped and not a bulleted dump of directive headings.
15. It has a cross-cutting-themes section that synthesizes patterns recurring across multiple utilities and ties each theme to the utilities it affects (real synthesis across Decisions, not a single restated directive).
16. It has an outliers section calling out where one or two utilities clearly diverge from the fleet, with the specific reason they stand out.
17. It has a recommendations section about what the reader's own maintenance program should tighten before the next review, connected to the findings.
18. It includes a sources/appendix section listing the Decisions and their links.

CHARTS:
19. At least 2 of the 3 expected ledger-derived charts are present (chart_findings_by_topic.png, chart_findings_by_utility.png, chart_directive_themes.png).

DECK (briefing_deck.pptx):
20. The deck exists with at least 8 slides covering what the Decisions cover, per-utility headlines, cross-cutting themes, outliers, and recommended actions.

CONSISTENCY ACROSS DELIVERABLES:
21. The briefing and deck are consistent with the SUBSTANTIVE findings — the utilities and themes they emphasize match the utilities and topics most represented among genuine findings, and the briefing's claims are grounded in those paraphrased findings (it does not lean on a utility with essentially no substantive findings, or quietly contradict the ledger's counts/themes).

Respond with ONLY this JSON object (no markdown fences, no extra text):
{"checks":[{"id":1,"pass":true,"reason":"..."},{"id":2,"pass":false,"reason":"..."},...all 21],"total_passed":<int>}
"""


def evaluate(ledger_text, briefing_text, briefing_words, deck_text, deck_slides, charts_present, extra_pngs):
    api_key = os.environ.get("FIREWORKS_API_KEY", "")
    if not api_key:
        return set(), [], "No FIREWORKS_API_KEY in environment"

    try:
        from openai import OpenAI
    except ImportError:
        return set(), [], "Could not import openai — pip install openai"

    client = OpenAI(api_key=api_key, base_url="https://api.fireworks.ai/inference/v1")

    combined = (
        "--- cross_utility_findings.csv (ledger summary) ---\n" + ledger_text
        + "\n\n--- reliability_briefing.docx ---\n"
        + f"(word count: {briefing_words})\n" + (briefing_text or "[not found or empty]")
        + "\n\n--- briefing_deck.pptx ---\n"
        + f"(slide count: {deck_slides})\n" + (deck_text or "[not found or empty]")
        + "\n\n--- charts present under /logs/agent/ ---\n"
        + f"Expected charts present: {charts_present}\n"
        + f"Other PNG files present: {extra_pngs}"
    )

    last_err = None
    for attempt in range(3):
        try:
            completion = client.chat.completions.create(
                model="accounts/fireworks/models/qwen3p7-plus",
                messages=[
                    {"role": "system", "content": EVALUATION_PROMPT},
                    {"role": "user", "content": combined},
                ],
                temperature=0,
            )
            body = (completion.choices[0].message.content or "").strip()
            match = re.search(r"\{[\s\S]*\}", body)
            if not match:
                raise ValueError(f"LLM did not return JSON: {body[:200]}")

            parsed = json.loads(match.group(0))
            items = parsed.get("checks", [])

            passed = set()
            for item in items:
                try:
                    cid = int(item["id"])
                except (KeyError, TypeError, ValueError):
                    continue
                if 1 <= cid <= LLM_CHECKS and item.get("pass"):
                    passed.add(cid)

            lines = []
            for item in sorted(items, key=lambda x: x.get("id", 0)):
                cid = item.get("id", "?")
                if isinstance(cid, int) and not (1 <= cid <= LLM_CHECKS):
                    continue
                tag = "PASS" if item.get("pass") else "FAIL"
                lines.append(f"  [{tag}] {cid}: {item.get('reason','')}")
            return passed, lines, None

        except Exception as exc:
            last_err = exc
            if attempt < 2:
                import time
                time.sleep(3 * (attempt + 1))

    return set(), [], f"All LLM attempts failed. Last error: {last_err}"


def write_reward(score, error=None):
    REWARD_PATH.parent.mkdir(parents=True, exist_ok=True)
    payload = {"reward": score}
    if error:
        payload["error"] = error
    with open(REWARD_PATH, "w") as f:
        json.dump(payload, f)
    try:
        REWARD_TXT.write_text(str(score))
    except OSError:
        pass


def main():
    ledger_rows = load_ledger()
    ledger_text = summarize_ledger(ledger_rows)
    briefing_text, briefing_words = load_docx(AGENT_DIR / "reliability_briefing.docx")
    deck_text, deck_slides = load_pptx(AGENT_DIR / "briefing_deck.pptx")
    charts_present, extra_pngs = list_charts()

    # Checks 1-21: run the LLM judge several times INDEPENDENTLY and majority-vote
    # each check, so no single judge call decides the majority of the reward.
    run_passes = []
    run_lines = []
    judge_errors = []
    for _ in range(NUM_JUDGE_RUNS):
        passed, lines, err = evaluate(
            ledger_text, briefing_text, briefing_words, deck_text, deck_slides, charts_present, extra_pngs
        )
        if err is None:
            run_passes.append(passed)
            run_lines = lines
        else:
            judge_errors.append(err)

    # Fail-closed: this is an llm-judge task, so a run must NOT earn a non-zero
    # reward from the deterministic depth checks alone. If every judge call failed,
    # write an explicit error state and score 0.0 (the depth checks are not counted).
    if not run_passes:
        err_msg = "llm_judge_unavailable: " + "; ".join(judge_errors[:3])
        detail = (
            f"Result: 0/{NUM_CHECKS} checks passed  reward=0.0  "
            f"(quality 1-21: 0/{LLM_CHECKS}, depth 22-30: not scored)\n"
            f"  [LLM ERROR] checks 1-21 could not be scored: {err_msg}\n"
            f"  [FAIL-CLOSED] depth checks 22-30 are not counted because the LLM judge did not run."
        )
        write_reward(0.0, error=err_msg)
        print(detail)
        try:
            (AGENT_DIR / "judge_justification.txt").write_text(detail)
        except OSError:
            pass
        return

    # Majority vote across the successful judge runs.
    vote = Counter()
    for passed in run_passes:
        vote.update(passed)
    threshold = len(run_passes) / 2.0
    llm_passed = {cid for cid, count in vote.items() if count > threshold}
    llm_lines = run_lines

    # Checks 22-30: deterministic depth-of-coverage from the ledger. These are
    # counted ONLY because the LLM judge above succeeded; they never produce a
    # non-zero reward on their own (see the fail-closed branch).
    depth_items, metrics = depth_checks(ledger_rows)
    depth_passed = sum(1 for it in depth_items if it["pass"])

    total_passed = len(llm_passed) + depth_passed
    score = round(total_passed / NUM_CHECKS, 4)

    lines = [
        f"Result: {total_passed}/{NUM_CHECKS} checks passed  reward={score}  "
        f"(quality 1-21: {len(llm_passed)}/{LLM_CHECKS} [majority of {len(run_passes)} judge runs], "
        f"depth 22-30: {depth_passed}/{DEPTH_CHECKS})"
    ]
    lines.extend(llm_lines)
    lines.append("  --- depth-of-coverage (deterministic, from ledger) ---")
    for it in depth_items:
        tag = "PASS" if it["pass"] else "FAIL"
        lines.append(f"  [{tag}] {it['id']}: {it['reason']}")
    lines.append("  depth metrics: " + json.dumps(metrics))
    detail = "\n".join(lines)

    write_reward(score)
    print(detail)
    try:
        (AGENT_DIR / "judge_justification.txt").write_text(detail)
    except OSError:
        pass


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        traceback.print_exc()
        write_reward(0.0, error=f"verifier_exception: {type(exc).__name__}: {exc}")

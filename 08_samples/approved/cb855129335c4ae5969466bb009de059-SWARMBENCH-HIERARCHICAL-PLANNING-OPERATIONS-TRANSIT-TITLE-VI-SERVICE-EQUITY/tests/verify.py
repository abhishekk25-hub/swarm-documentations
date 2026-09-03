#!/usr/bin/env python3
"""Deterministic executable verifier for the FY2026 Title VI service-equity task.

The deliverable is a Board packet with three parts:
  * /logs/agent/output.xlsx       -- the structured analysis (route matrix,
                                     system rollup, conflict/exclusion logs,
                                     source register, comment dispositions,
                                     evidence register, chart manifest, summary)
  * /logs/agent/board_report.md   -- the long-form, evidence-backed written
                                     report (per-route justifications, transit-
                                     dependent context, public-comment synthesis)
  * three PNG charts in /logs/agent

Scoring is a transparent WEIGHTED RUBRIC. Each category is scored as its own
earned/possible ratio with smooth partial credit, then combined with fixed,
mode-agnostic weights. There is no per-mode branch, cap, floor or multiplier.

The numeric backbone is RECOMPUTED from the raw input tables in
tests/reference_inputs/ (no pre-authored answer file). The browsing check
compares the agent's submitted citations against acceptable source domains and
load-bearing fact tokens (the verifier never accesses the network). The writing
checks are coverage/consistency based (per-route justification present and
quantified, every public comment dispositioned, required sections present and
internally consistent with the recomputed numbers).
"""

import argparse
import csv
import json
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REF_DIR = os.path.join(HERE, "reference_inputs")

DATE_RE = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
URL_RE = re.compile(r"https?://[^\s)\]]+", re.IGNORECASE)

# Mode-agnostic rubric weights (sum to 1.0). They reflect the relative size of
# each part of the Board deliverable, not any single/multi distinction.
WEIGHTS = {
    "numeric_backbone": 0.30,
    "source_register": 0.22,
    "route_justifications": 0.18,
    "comment_synthesis": 0.15,
    "report_structure": 0.10,
    "charts": 0.05,
}

REQUIRED_SHEETS = [
    "route_equity_matrix", "system_rollup", "conflict_log", "exclusion_log",
    "source_register", "comment_disposition", "evidence_register",
    "chart_manifest", "summary",
]

REPORT_SECTIONS = [
    ["executive summary"],
    ["background", "package description", "package overview"],
    ["methodology"],
    ["major service change"],
    ["aggregate", "disparate impact", "equity finding"],
    ["route-by-route", "route by route", "per-route", "route-level"],
    ["transit-dependent", "transit dependent", "zero-vehicle", "limited english", "older adult"],
    ["data limitation", "excluded route", "exclusion"],
    ["reconciliation", "schedule conflict", "source of record"],
    ["public comment"],
    ["mitigation", "disposition", "recommendation"],
    ["regulatory basis", "citations", "authorities"],
]

MEMO_DISCLOSURES = ["disparate impact", "disproportionate burden",
                    "people-trips", "substantial legitimate justification"]
DISPOSITION_KEYWORDS = [
    "incorporat", "mitigat", "noted", "declin", "defer", "refer", "acknowledg",
    "out of scope", "out-of-scope", "adopt", "retain", "restore", "no change",
    "decline", "address", "response", "respond",
]
BANNED_TOKENS = ("lorem", "[insert", "tbd", "todo", "placeholder", "xxxx", "<fill")

REQUIRED_CHART_KEYS = [
    ["people", "trip"],     # people-trip change by route
    ["share"],              # burden share vs system share
    ["major", "sub"],       # major vs sub-threshold contribution
]

REL_TOL = 0.01
SHARE_TOL = 0.01
RATIO_TOL = 0.02
JUSTIFY_MIN_CHARS = 160
REPORT_MIN_CHARS = 9000


def norm(v):
    return "" if v is None else str(v).strip()


def num(v):
    s = norm(v).replace(",", "").replace("$", "").replace("%", "").strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def yesno(v):
    s = norm(v).lower()
    if s in ("yes", "y", "true", "1", "triggered", "found"):
        return True
    if s in ("no", "n", "false", "0", "not triggered", "none", "not found", "cleared"):
        return False
    return None


def close_rel(a, b, rel=REL_TOL):
    if a is None:
        return False
    if b == 0:
        return abs(a) <= 1.0
    return abs(a - b) <= abs(b) * rel


def close_abs(a, b, tol):
    return a is not None and abs(a - b) <= tol


# ---------------------------------------------------------------------------
# Recompute the numeric reference from the raw input tables.
# ---------------------------------------------------------------------------
def load_csv(name):
    with open(os.path.join(REF_DIR, name)) as f:
        return list(csv.DictReader(f))


def build_reference():
    pc = json.load(open(os.path.join(REF_DIR, "policy_constants.json")))
    pkg = {r["route_id"]: r for r in load_csv("service_change_package_fy2026.csv")}
    gtfs = {r["route_id"]: r for r in load_csv("gtfs_schedule_export.csv")}
    demo = {r["route_id"]: r for r in load_csv("walkshed_demographics_acs.csv")}
    cal = {r["day_type"]: int(r["service_days_per_year"]) for r in load_csv("service_calendar_fy2026.csv")}

    major_pct = pc["major_change_pct"]
    ratio_thr = pc["di_db_ratio_threshold"]
    sys_min = pc["system_minority_share"]
    sys_low = pc["system_low_income_share"]
    overlap_min = pc["acs_overlap_min_pct"]

    def pct_change(b, a):
        if b == 0:
            return 1.0 if a != 0 else 0.0
        return abs(a - b) / b

    routes = {}
    for rid, p in pkg.items():
        d = demo[rid]
        days = cal.get(p["affected_day_type"], 0)
        change_type = p["change_type"]
        drm_b = float(p["daily_revenue_miles_before"]); drm_a = float(p["daily_revenue_miles_after"])
        trips_b = int(p["daily_trips_before"]); trips_a = int(p["daily_trips_after"])
        gtfs_trips = int(gtfs[rid]["scheduled_daily_trips_post_change"])
        conflict = (gtfs_trips != trips_a)
        pop_b = int(d["walkshed_population_before"]); pop_a = int(d["walkshed_population_after"])
        mshare = float(d["minority_population_share"]); lshare = float(d["low_income_population_share"])
        overlap = float(d["acs_blockgroup_overlap_pct"])

        atrips_b = trips_b * days; atrips_a = trips_a * days
        d_ov = pop_a * atrips_a - pop_b * atrips_b
        d_mi = pop_a * mshare * atrips_a - pop_b * mshare * atrips_b
        d_lo = pop_a * lshare * atrips_a - pop_b * lshare * atrips_b

        is_major = max(pct_change(drm_b, drm_a), pct_change(trips_b, trips_a)) >= major_pct
        has_change = change_type != "none"
        included = has_change and overlap >= overlap_min

        routes[rid] = dict(
            change_type=change_type, is_major=is_major, has_change=has_change,
            included=included, overlap=overlap, conflict=conflict,
            governing_trips_after=trips_a, gtfs_trips=gtfs_trips,
            d_overall=d_ov, d_minority=d_mi, d_low_income=d_lo,
            annual_trips_before=atrips_b, annual_trips_after=atrips_a,
        )

    incl = [r for r in routes.values() if r["included"]]
    D_ov = sum(r["d_overall"] for r in incl)
    D_mi = sum(r["d_minority"] for r in incl)
    D_lo = sum(r["d_low_income"] for r in incl)
    min_share = D_mi / D_ov if D_ov else 0.0
    low_share = D_lo / D_ov if D_ov else 0.0
    di_ratio = min_share / sys_min if sys_min else 0.0
    db_ratio = low_share / sys_low if sys_low else 0.0

    return dict(
        routes=routes,
        conflict_ids=[rid for rid, r in routes.items() if r["conflict"]],
        excluded_ids=[rid for rid, r in routes.items() if r["has_change"] and not r["included"]],
        included_ids=[rid for rid, r in routes.items() if r["included"]],
        major_ids=[rid for rid, r in routes.items() if r["has_change"] and r["is_major"]],
        changed_ids=[rid for rid, r in routes.items() if r["has_change"]],
        D_overall=D_ov, D_minority=D_mi, D_low_income=D_lo,
        minority_share_of_change=min_share, low_income_share_of_change=low_share,
        di_ratio=di_ratio, db_ratio=db_ratio,
        disparate_impact=di_ratio >= ratio_thr, disproportionate_burden=db_ratio >= ratio_thr,
        system_minority_share=sys_min, system_low_income_share=sys_low, ratio_threshold=ratio_thr,
    )


def load_sources():
    with open(os.path.join(REF_DIR, "sources_reference.json")) as f:
        return json.load(f)["sources"]


def load_comment_ids():
    ids = []
    with open(os.path.join(REF_DIR, "public_comment_log.txt")) as f:
        for line in f:
            m = re.match(r"\s*(PC-\d+)\b", line)
            if m:
                ids.append(m.group(1))
    return ids


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------
def read_sheet(wb, name):
    if name not in wb.sheetnames:
        return None
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    header = [norm(h).lower() for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(c is None or norm(c) == "" for c in raw):
            continue
        out.append({header[i]: (raw[i] if i < len(raw) else None) for i in range(len(header))})
    return out


def kv_sheet(wb, name):
    rows = read_sheet(wb, name) or []
    out = {}
    for r in rows:
        keys = list(r.keys())
        if len(keys) >= 2:
            k = norm(r[keys[0]]).lower()
            if k:
                out[k] = r[keys[1]]
    return out


def get(row, *names):
    for n in names:
        if n in row and norm(row[n]) != "":
            return row[n]
    return None


def is_png(path):
    try:
        with open(path, "rb") as f:
            sig = f.read(8)
        return sig.startswith(b"\x89PNG\r\n\x1a\n") and os.path.getsize(path) > 3000
    except OSError:
        return False


# ---------------------------------------------------------------------------
# Category scorers -- each returns (earned, possible, lines)
# ---------------------------------------------------------------------------
def score_numeric_backbone(ref, wb):
    R = ref["routes"]
    matrix = read_sheet(wb, "route_equity_matrix") or []
    out_by_id = {}
    for r in matrix:
        rid = norm(get(r, "route_id"))
        if rid and rid not in out_by_id:
            out_by_id[rid] = r
    rollup = kv_sheet(wb, "system_rollup")
    summary = kv_sheet(wb, "summary")
    conflicts = read_sheet(wb, "conflict_log") or []
    exclusions = read_sheet(wb, "exclusion_log") or []

    e = p = 0
    lines = []

    def per_route(label, fn):
        nonlocal e, p
        earned = sum(1 for rid in R if rid in out_by_id and fn(rid, out_by_id[rid]))
        e += earned; p += len(R)
        lines.append(f"    {earned:>3}/{len(R):<3} route: {label}")

    per_route("major-change screen", lambda rid, row: yesno(get(row, "is_major_change", "major_change")) == R[rid]["is_major"])
    per_route("rollup inclusion", lambda rid, row: yesno(get(row, "included_in_equity_rollup", "included", "in_rollup")) == R[rid]["included"])
    per_route("governing post-change trips", lambda rid, row: close_abs(num(get(row, "daily_trips_after_governing", "daily_trips_after")), R[rid]["governing_trips_after"], 0.5))
    per_route("people-trips delta", lambda rid, row: close_rel(num(get(row, "people_trips_delta", "people_trip_delta")), R[rid]["d_overall"]))
    per_route("minority people-trips delta", lambda rid, row: close_rel(num(get(row, "minority_people_trips_delta", "minority_delta")), R[rid]["d_minority"]))
    per_route("low-income people-trips delta", lambda rid, row: close_rel(num(get(row, "low_income_people_trips_delta", "low_income_delta")), R[rid]["d_low_income"]))

    matrix_present = min(len(out_by_id), len(R))
    e += matrix_present; p += len(R)
    lines.append(f"    {matrix_present:>3}/{len(R):<3} matrix: every package route present")

    def rnum(*names):
        for n in names:
            if n in rollup:
                return num(rollup[n])
        return None

    roll_checks = [
        ("D_overall", close_rel(rnum("d_overall", "total_people_trips_change"), ref["D_overall"])),
        ("D_minority", close_rel(rnum("d_minority", "minority_people_trips_change"), ref["D_minority"])),
        ("D_low_income", close_rel(rnum("d_low_income", "low_income_people_trips_change"), ref["D_low_income"])),
        ("minority share", close_abs(rnum("minority_share_of_change", "minority_burden_share"), ref["minority_share_of_change"], SHARE_TOL)),
        ("low-income share", close_abs(rnum("low_income_share_of_change", "low_income_burden_share"), ref["low_income_share_of_change"], SHARE_TOL)),
        ("di_ratio", close_abs(rnum("di_ratio", "disparate_impact_ratio"), ref["di_ratio"], RATIO_TOL)),
        ("db_ratio", close_abs(rnum("db_ratio", "disproportionate_burden_ratio"), ref["db_ratio"], RATIO_TOL)),
    ]
    rok = sum(1 for _, ok in roll_checks if ok)
    e += rok; p += len(roll_checks)
    lines.append(f"    {rok:>3}/{len(roll_checks):<3} system rollup metrics correct")

    di = yesno(rollup.get("disparate_impact_finding") or rollup.get("disparate_impact")
               or summary.get("disparate_impact_finding") or summary.get("disparate_impact"))
    db = yesno(rollup.get("disproportionate_burden_finding") or rollup.get("disproportionate_burden")
               or summary.get("disproportionate_burden_finding") or summary.get("disproportionate_burden"))
    find = [di is not None, di == ref["disparate_impact"], db is not None, db == ref["disproportionate_burden"]]
    fok = sum(1 for x in find if x)
    e += fok; p += 4
    lines.append(f"    {fok:>3}/{4:<3} headline findings stated & correct (DI triggers, DB clears)")

    cids = set(ref["conflict_ids"])
    conf_seen = {norm(get(c, "route_id")) for c in conflicts if norm(get(c, "route_id"))}
    conf_ok = 0
    seen_conf_rids = set()
    for c in conflicts:
        rid = norm(get(c, "route_id"))
        if rid in seen_conf_rids:
            continue
        seen_conf_rids.add(rid)
        if rid in cids and close_abs(num(get(c, "resolved_value", "resolved", "governing_value")), R[rid]["governing_trips_after"], 0.5):
            conf_ok += 1
    cscore = len(conf_seen & cids) + conf_ok
    e += cscore; p += 2 * len(cids)
    lines.append(f"    {cscore:>3}/{2 * len(cids):<3} conflicts identified & resolved to governing value")

    xids = set(ref["excluded_ids"])
    xseen = {norm(get(x, "route_id")) for x in exclusions if norm(get(x, "route_id"))}
    xkept = sum(1 for rid in xids if rid in out_by_id and yesno(get(out_by_id[rid], "included_in_equity_rollup", "included")) is False)
    xscore = len(xseen & xids) + xkept
    e += xscore; p += 2 * len(xids)
    lines.append(f"    {xscore:>3}/{2 * len(xids):<3} data-gap routes excluded & kept out of rollup")

    return e, p, lines, out_by_id


def score_source_register(sources, wb):
    rows = (read_sheet(wb, "source_register") or []) + (read_sheet(wb, "regulatory_basis") or [])
    e = 0; p = 3 * len(sources)
    hit = 0
    for s in sources:
        toks = s["fact_tokens"]; doms = s["domains"]
        fact_row = None
        for row in rows:
            blob = " ".join(norm(v).lower() for v in row.values())
            if any(all(t in blob for t in grp) for grp in toks):
                fact_row = row
                break
        if fact_row is None:
            continue
        blob = " ".join(norm(v).lower() for v in fact_row.values())
        urls = URL_RE.findall(blob)
        dom_ok = any(any(d in u.lower() for d in doms) for u in urls)
        date_ok = bool(DATE_RE.search(blob))
        sc = 1 + (1 if dom_ok else 0) + (1 if date_ok else 0)
        e += sc
        if sc == 3:
            hit += 1
    lines = [f"    {e:>3}/{p:<3} authorities verified (fact token + official-domain URL + frozen date)",
             f"        ({hit}/{len(sources)} fully verified with all three)"]
    return e, p, lines


def score_route_justifications(ref, report):
    changed = ref["changed_ids"]
    e = 0; p = len(changed)
    paras = re.split(r"\n\s*\n", report)
    for rid in changed:
        ok = False
        for para in paras:
            if rid in para and len(para.strip()) >= JUSTIFY_MIN_CHARS and re.search(r"\d", para):
                ok = True
                break
        if not ok:  # fall back: route id on a line with a number and enough text nearby
            for m in re.finditer(re.escape(rid), report):
                window = report[max(0, m.start() - 30): m.start() + JUSTIFY_MIN_CHARS + 60]
                if len(window) >= JUSTIFY_MIN_CHARS and re.search(r"\d", window.replace(rid, "")):
                    ok = True
                    break
        e += 1 if ok else 0
    lines = [f"    {e:>3}/{p:<3} changed routes with a quantified narrative justification in the report"]
    return e, p, lines


def score_comment_synthesis(comment_ids, report, wb):
    disp_rows = read_sheet(wb, "comment_disposition") or []
    by_id = {}
    for r in disp_rows:
        cid = norm(get(r, "comment_id", "id", "pc_id")).upper()
        if cid:
            by_id[cid] = r
    rlow = report.lower()
    e = 0; p = len(comment_ids)
    for cid in comment_ids:
        ok = False
        row = by_id.get(cid)
        if row:
            blob = " ".join(norm(v).lower() for v in row.values())
            disp = norm(get(row, "disposition", "response", "response_note", "resolution", "action"))
            if disp or any(k in blob for k in DISPOSITION_KEYWORDS):
                ok = True
        if not ok:
            i = rlow.find(cid.lower())
            if i >= 0:
                window = rlow[i: i + 320]
                if any(k in window for k in DISPOSITION_KEYWORDS):
                    ok = True
        e += 1 if ok else 0
    lines = [f"    {e:>3}/{p:<3} public comments dispositioned (sheet or report synthesis)"]
    return e, p, lines


def score_report_structure(ref, report):
    rlow = report.lower()
    e = 0; p = 0
    lines = []

    sect_hits = sum(1 for grp in REPORT_SECTIONS if any(k in rlow for k in grp))
    e += sect_hits; p += len(REPORT_SECTIONS)
    lines.append(f"    {sect_hits:>3}/{len(REPORT_SECTIONS):<3} required report sections present")

    disc = sum(1 for d in MEMO_DISCLOSURES if d in rlow)
    e += disc; p += len(MEMO_DISCLOSURES)
    lines.append(f"    {disc:>3}/{len(MEMO_DISCLOSURES):<3} key equity disclosures present")

    di_str = f"{ref['di_ratio']:.2f}"
    consistency = 0
    if di_str in rlow or f"{ref['di_ratio']:.3f}" in rlow:
        consistency += 1
    if ("disparate impact" in rlow) and any(w in rlow for w in ("found", "triggered", "is a disparate")):
        consistency += 1
    e += consistency; p += 2
    lines.append(f"    {consistency:>3}/{2:<3} report numbers consistent with recomputed finding")

    no_ph = 1 if (rlow and not any(b in rlow for b in BANNED_TOKENS)) else 0
    e += no_ph; p += 1
    lines.append(f"    {no_ph:>3}/{1:<3} no placeholder text")

    length_ok = 1 if len(report) >= REPORT_MIN_CHARS else round(len(report) / REPORT_MIN_CHARS, 2)
    length_ok = min(length_ok, 1)
    e += length_ok; p += 1
    lines.append(f"    {round(length_ok,2):>3}/{1:<3} report length (>= {REPORT_MIN_CHARS} chars; got {len(report)})")

    return e, p, lines


def score_charts(wb, agent_dir):
    charts = read_sheet(wb, "chart_manifest") or []
    e = 0; p = len(REQUIRED_CHART_KEYS)
    for keys in REQUIRED_CHART_KEYS:
        for c in charts:
            blob = " ".join(norm(v).lower() for v in c.values())
            fname = None
            for k, v in c.items():
                if ("file" in k or "name" in k or "path" in k) and norm(v).lower().endswith(".png"):
                    fname = norm(v)
            if fname and all(kw in blob for kw in keys):
                path = fname if os.path.isabs(fname) else os.path.join(agent_dir, fname)
                if is_png(path):
                    e += 1
                    break
    lines = [f"    {e:>3}/{p:<3} required PNG charts produced + manifested"]
    return e, p, lines


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--agent-output", default="/logs/agent/output.xlsx")
    ap.add_argument("--agent-dir", default="/logs/agent")
    ap.add_argument("--report", default="/logs/agent/board_report.md")
    ap.add_argument("--reward-out", default="/logs/verifier/reward.txt")
    args = ap.parse_args()

    def finish(score, lines):
        score = round(min(float(score), 1.0), 4)
        try:
            os.makedirs(os.path.dirname(args.reward_out), exist_ok=True)
            with open(args.reward_out, "w") as f:
                f.write(str(score))
        except OSError as ex:
            print(f"WARN: could not write reward: {ex}")
        try:
            with open(os.path.join(args.agent_dir, "judge_justification.txt"), "w") as f:
                f.write(f"Score: {score}\n\n" + "\n".join(lines))
        except OSError:
            pass
        print(f"Score: {score}")
        print("\n".join(lines))

    ref = build_reference()
    sources = load_sources()
    comment_ids = load_comment_ids()

    try:
        wb = openpyxl.load_workbook(args.agent_output, data_only=True, read_only=True)
    except Exception as ex:
        wb = None
        wb_err = str(ex)
    else:
        wb_err = None

    report = ""
    if os.path.exists(args.report):
        try:
            with open(args.report, encoding="utf-8", errors="ignore") as f:
                report = f.read()
        except OSError:
            report = ""

    cats = {}
    detail = []

    if wb is not None:
        e, p, l, _ = score_numeric_backbone(ref, wb)
        cats["numeric_backbone"] = (e, p); detail.append(("numeric_backbone", l))
        e, p, l = score_source_register(sources, wb)
        cats["source_register"] = (e, p); detail.append(("source_register", l))
        e, p, l = score_comment_synthesis(comment_ids, report, wb)
        cats["comment_synthesis"] = (e, p); detail.append(("comment_synthesis", l))
        e, p, l = score_charts(wb, args.agent_dir)
        cats["charts"] = (e, p); detail.append(("charts", l))
    else:
        for k in ("numeric_backbone", "source_register", "comment_synthesis", "charts"):
            cats[k] = (0, 1)
        detail.append(("workbook", [f"    no usable workbook at {args.agent_output}: {wb_err}"]))

    e, p, l = score_route_justifications(ref, report)
    cats["route_justifications"] = (e, p); detail.append(("route_justifications", l))
    e, p, l = score_report_structure(ref, report)
    cats["report_structure"] = (e, p); detail.append(("report_structure", l))

    score = 0.0
    cat_lines = []
    for cat, w in WEIGHTS.items():
        e, p = cats.get(cat, (0, 1))
        ratio = (e / p) if p else 0.0
        score += w * ratio
        cat_lines.append(f"  [{w:.2f}] {cat:22} {round(ratio,4):>6}  ({round(e,2)}/{p})")

    lines = [
        f"Title VI service-equity weighted score = {round(score,4)}",
        f"reference: disparate_impact={ref['disparate_impact']} (di_ratio={round(ref['di_ratio'],4)}),"
        f" disproportionate_burden={ref['disproportionate_burden']} (db_ratio={round(ref['db_ratio'],4)})",
        f"workbook sheets present: {[s for s in REQUIRED_SHEETS if wb is not None and s in wb.sheetnames]}",
        "",
        "category weighted ratios:",
    ] + cat_lines + ["", "detail:"]
    for _cat, l in detail:
        lines.extend(l)
    finish(score, lines)


if __name__ == "__main__":
    main()

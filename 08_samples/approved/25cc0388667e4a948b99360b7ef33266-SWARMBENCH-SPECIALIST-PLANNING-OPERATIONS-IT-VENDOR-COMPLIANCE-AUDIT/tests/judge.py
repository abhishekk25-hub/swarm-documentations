#!/usr/bin/env python3
"""
Verifier for the IT Vendor Compliance Audit task.
Sends all three agent deliverables to an LLM judge that scores
25 independent boolean criteria. Final reward = passed / 25.
"""

import json
import os
import re
import traceback
from pathlib import Path

AGENT_DIR = Path("/logs/agent")
REWARD_PATH = Path("/logs/verifier/reward.json")
NUM_CHECKS = 25


def load_xlsx_as_text(fpath):
    if not fpath.exists():
        return "[output.xlsx not found]"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(fpath), data_only=True, read_only=True)
    except Exception as exc:
        return f"[Could not open output.xlsx: {exc}]"

    parts = [f"Sheets in workbook: {', '.join(wb.sheetnames)}"]
    for sname in wb.sheetnames:
        ws = wb[sname]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            parts.append(f"\n[{sname}] — empty sheet")
            continue
        hdr = [str(c or "") for c in all_rows[0]]
        parts.append(f"\n[{sname}] {len(all_rows)-1} data rows")
        parts.append("Columns: " + " | ".join(hdr))
        for idx, row in enumerate(all_rows[1:], 1):
            parts.append(f"  {idx}. " + " | ".join(str(c or "") for c in row))
    return "\n".join(parts)


def load_evidence_summary(fpath):
    if not fpath.exists():
        return "[evidence_log.json not found]"
    try:
        raw = fpath.read_text(encoding="utf-8", errors="replace")
        entries = json.loads(raw)
    except Exception as exc:
        return f"[Could not parse evidence_log.json: {exc}]"
    if not isinstance(entries, list):
        return "[evidence_log.json is not a JSON array]"
    n = len(entries)
    if n <= 40:
        shown = entries
        label = f"Total entries: {n}. All shown:"
    else:
        first = entries[:15]
        mid_start = n // 2 - 5
        middle = entries[mid_start:mid_start + 10]
        last = entries[-15:]
        shown = first + middle + last
        label = (f"Total entries: {n}. Showing first 15, middle 10, last 15 "
                 f"(entries 1-15, {mid_start+1}-{mid_start+10}, {n-14}-{n}):")
    return label + "\n" + json.dumps(shown, indent=2, ensure_ascii=False)


def load_report(fpath):
    if not fpath.exists():
        return ""
    try:
        return fpath.read_text(encoding="utf-8", errors="replace")[:55000]
    except OSError:
        return ""


EVALUATION_PROMPT = """You are a strict but fair evaluator for a vendor compliance audit task.

An agent was asked to audit 25 SaaS vendors — Slack, Zoom, Atlassian, Notion,
HubSpot, Dropbox, Salesforce, ServiceNow, Okta, DocuSign, Box, Monday.com,
Workday, Zendesk, Twilio, Snowflake, Datadog, Cloudflare, PagerDuty, Elastic,
Confluent, MongoDB, GitLab, Airtable, and Figma — across SLA commitments,
security certifications, data privacy, and incident history. The agent should
have produced three deliverables under /logs/agent/: output.xlsx,
evidence_log.json, and risk_report.md.

Score the output on these 25 independent boolean checks. Mark each PASS or FAIL.

COMPLIANCE MATRIX (output.xlsx):
 1. The workbook has a sheet named vendor_audit with data rows.
 2. vendor_audit has one row per vendor (25 total) with at least vendor_name, sla_claim, cert_registry_status, and overall_risk columns.
 3. The cert_registry_status column shows realistic variation — at least 3 vendors show 'Not found', 'Expired', or 'Unverified'. Among 25 SaaS vendors, not all hold FedRAMP, CSA STAR, and ISO 27001 simultaneously, so genuine registry checking will find multiple vendors whose claims cannot be fully confirmed.
 4. At least 3 vendors have overall_risk = 'High'. A 25-vendor audit spanning two years of activity should surface some vendors with serious compliance concerns worth escalating.
 5. A risk_ranking sheet exists with 25 rows containing vendor_name, risk_rank, risk_level, and justification.
 6. The top 3 vendors by risk_rank (ranks 1, 2, 3) all have risk_level = 'High', and risk_rank runs sequentially 1 through 25 without gaps or repeats.
 7. At least 12 vendors have a sla_evidence_url that references an actual status or availability page — identifiable by containing 'status', 'statuspage', 'uptime', or 'availability' in the URL or subdomain — rather than just a generic trust, compliance, or homepage URL. The instruction explicitly requires verifying the uptime claim against status page history, not just recording the vendor's marketing claim.
 8. At least 20 vendors list specific certifications in cert_claimed (SOC 2, ISO 27001, FedRAMP, etc.).
 9. At least 10 vendors have an incident_summary that cites a specific month and year (e.g. 'March 2025', 'January 2026') rather than just saying 'None found' or giving a vague description without dates.
10. The risk_ranking justification field for the top 5 vendors (risk_rank 1 through 5) each mentions findings from at least two different compliance domains — for example, referencing both a certification gap and an incident, or both an SLA concern and a privacy finding. A risk ranking synthesized from a 4-domain audit should reflect multi-domain evidence, not a single data point.

EVIDENCE LOG (evidence_log.json):
You are shown a sample from the beginning, middle, and end of the file.
Use this sample to judge the quality of the full file.
11. The file contains at least 100 entries.
12. Every shown entry has all five required fields: vendor_name, domain, claim_made, verification_url, and verbatim_excerpt.
13. Every domain value in shown entries is one of: sla, certs, privacy, incidents.
14. Shown entries reference at least 20 of the 25 vendors by name.
15. Shown entries cover all four domains: sla, certs, privacy, and incidents.
16. Among the last 15 shown entries, at least 10 have a verbatim_excerpt that is a genuine sentence of 20 or more words taken from a source page — not "No excerpt available", not a stub, and not obviously templated text repeated across vendors.
17. Among the last 15 shown entries, at least 10 have a verification_url that is a specific, vendor-appropriate URL starting with https:// pointing to a subpage like a trust center, status page, or privacy policy — not just the vendor's homepage root (e.g. https://slack.com alone would not count, but https://slack.com/trust/compliance would).
18. The verbatim_excerpt values across the shown entries are genuinely distinct from each other — not the same sentence template repeated with different vendor names swapped in. Each excerpt should read like it was pulled from a different source page.

RISK REPORT (risk_report.md):
19. The report exists and is between 2000 and 3000 words.
20. It opens with an executive summary covering the overall compliance picture across all 25 vendors.
21. It covers at least three of the four audit domains and names specific vendors in each section.
22. It explicitly assigns a 'High' risk designation to at least 3 vendors and gives concrete supporting evidence for each — not just ranking them highest, but naming the specific risk drivers.
23. It identifies at least one cross-vendor pattern, such as a certification gap affecting multiple vendors or a shared privacy weakness.
24. It closes with a contract renewal action plan that assigns a renewal decision of Renew, Renegotiate, or Escalate to Legal for at least three of the five highest-risk vendors, with specific contract clauses or compliance gaps cited as the reason for each decision.

CONSISTENCY:
25. The vendors the report calls out as highest risk match the top 5 in the risk_ranking sheet.

Respond with ONLY this JSON object (no markdown fences, no extra text):
{"checks":[{"id":1,"pass":true,"reason":"..."},{"id":2,"pass":false,"reason":"..."},...all 25],"total_passed":<int>}
"""


def evaluate(xlsx_text, evidence_text, report_text):
    api_key = os.environ.get("FIREWORKS_API_KEY", "")
    if not api_key:
        return 0.0, "No FIREWORKS_API_KEY in environment"

    try:
        from openai import OpenAI
    except ImportError:
        return 0.0, "Could not import openai — pip install openai"

    client = OpenAI(api_key=api_key, base_url="https://api.fireworks.ai/inference/v1")

    combined = (
        "--- output.xlsx ---\n" + xlsx_text
        + "\n\n--- evidence_log.json ---\n" + evidence_text
        + "\n\n--- risk_report.md ---\n" + (report_text or "[not found or empty]")
    )

    last_err = None
    for attempt in range(3):
        try:
            completion = client.chat.completions.create(
                model="accounts/fireworks/models/qwen3p7-plus",
                messages=[
                    {"role": "system", "content": EVALUATION_PROMPT},
                    {"role": "user", "content": combined},
                ],
                temperature=0,
            )
            body = (completion.choices[0].message.content or "").strip()
            match = re.search(r"\{[\s\S]*\}", body)
            if not match:
                raise ValueError(f"LLM did not return JSON: {body[:200]}")

            parsed = json.loads(match.group(0))
            items = parsed.get("checks", [])

            passed = set()
            for item in items:
                try:
                    cid = int(item["id"])
                except (KeyError, TypeError, ValueError):
                    continue
                if 1 <= cid <= NUM_CHECKS and item.get("pass"):
                    passed.add(cid)

            score = len(passed) / NUM_CHECKS
            lines = [f"Result: {len(passed)}/{NUM_CHECKS}"]
            for item in sorted(items, key=lambda x: x.get("id", 0)):
                tag = "PASS" if item.get("pass") else "FAIL"
                lines.append(f"  [{tag}] {item.get('id','?')}: {item.get('reason','')}")

            return round(score, 4), "\n".join(lines)

        except Exception as exc:
            last_err = exc
            if attempt < 2:
                import time
                time.sleep(3 * (attempt + 1))

    return 0.0, f"All LLM attempts failed. Last error: {last_err}"


def main():
    xlsx_text = load_xlsx_as_text(AGENT_DIR / "output.xlsx")
    evidence_text = load_evidence_summary(AGENT_DIR / "evidence_log.json")
    report_text = load_report(AGENT_DIR / "risk_report.md")

    score, detail = evaluate(xlsx_text, evidence_text, report_text)

    REWARD_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(REWARD_PATH, "w") as f:
        json.dump({"reward": score}, f)

    print(detail)

    try:
        (AGENT_DIR / "judge_justification.txt").write_text(detail)
    except OSError:
        pass


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        REWARD_PATH.parent.mkdir(parents=True, exist_ok=True)
        with open(REWARD_PATH, "w") as f:
            json.dump({"reward": 0.0}, f)

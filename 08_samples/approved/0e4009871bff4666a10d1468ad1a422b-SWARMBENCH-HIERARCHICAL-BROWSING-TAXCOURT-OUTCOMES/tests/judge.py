#!/usr/bin/env python3
import argparse
import csv
import glob
import json
import os
import re
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
CORPUS = os.path.join(HERE, "corpus")
COHORT = os.path.join(HERE, "cohort.json")
API = "https://api.fireworks.ai/inference/v1/chat/completions"
MODEL = "accounts/fireworks/models/gpt-oss-120b"
W = {"relief": 0.30, "penalty": 0.25, "dispositive": 0.15, "issue": 0.10,
     "judge": 0.07, "years": 0.07, "amount": 0.06}
OPINIONS_WEIGHT = 0.88
DELIVERABLES_WEIGHT = 0.12
MAXWORKERS = 4
CAP_STOP = set("petitioner petitioners respondent commissioner deceased estate personal "
               "representative et al trustee inc llc llp co company and the of v".split())

SYS = ("You grade factual extractions from United States Tax Court memorandum opinions. "
       "You receive the full text of one opinion and the values a worker recorded for it. "
       "Decide each judgement strictly from the opinion text, then reply with only a JSON object.")


def norm(s):
    s = unicodedata.normalize("NFKC", str(s or ""))
    return re.sub(r"\s+", " ", s).strip().lower()


def cap_tokens(s):
    return {w for w in re.findall(r"[a-z0-9]+", norm(s)) if len(w) > 2 and w not in CAP_STOP}


def col(row, *keys):
    for k in row:
        if any(t in k for t in keys):
            return k
    return None


def _rows_from(table):
    if not table:
        return []
    hdr = [norm(h) for h in table[0]]
    out = []
    for raw in table[1:]:
        if raw is None or all(c is None or norm(c) == "" for c in raw):
            continue
        out.append({hdr[i]: ("" if i >= len(raw) or raw[i] is None else str(raw[i]).strip())
                    for i in range(len(hdr))})
    return out


def _workbook(agent_dir):
    xp = os.path.join(agent_dir, "opinions.xlsx")
    if os.path.isfile(xp):
        try:
            return openpyxl.load_workbook(xp, data_only=True, read_only=True)
        except Exception:
            return None
    return None


def read_table(agent_dir):
    wb = _workbook(agent_dir)
    if wb is not None:
        sheet = next((n for n in wb.sheetnames if "opinion" in norm(n)), wb.sheetnames[0])
        rows = _rows_from(list(wb[sheet].iter_rows(values_only=True)))
        if rows:
            return rows
    cands = ([os.path.join(agent_dir, n) for n in ("opinions.csv", "opinions.xlsx.csv")]
             + sorted(glob.glob(os.path.join(agent_dir, "*opinion*.csv"))))
    for p in cands:
        if os.path.isfile(p):
            try:
                with open(p, newline="", encoding="utf-8", errors="replace") as fh:
                    rows = _rows_from([r for r in csv.reader(fh)])
                if rows:
                    return rows
            except Exception:
                continue
    return []


def read_summary(agent_dir):
    wb = _workbook(agent_dir)
    if wb is not None:
        sn = next((n for n in wb.sheetnames if "summ" in norm(n)), None)
        if sn:
            rows = list(wb[sn].iter_rows(values_only=True))
            txt = "\n".join(" | ".join("" if c is None else str(c) for c in r) for r in rows if r)
            if txt.strip():
                return txt[:2000]
    return "(no summary sheet provided)"


def corpus_text(rid):
    p = os.path.join(CORPUS, rid + ".txt")
    try:
        t = open(p, encoding="utf-8", errors="replace").read()
    except OSError:
        return ""
    if len(t) > 120000:
        t = t[:80000] + "\n...\n" + t[-40000:]
    return t


def read_brief(agent_dir):
    for nm in ("brief.md", "memo.md", "report.md"):
        p = os.path.join(agent_dir, nm)
        if os.path.isfile(p):
            return open(p, encoding="utf-8", errors="ignore").read()
    return ""


def call_model(user, key):
    payload = {"model": MODEL, "temperature": 0, "max_tokens": 6000,
               "messages": [{"role": "system", "content": SYS}, {"role": "user", "content": user}]}
    last = ""
    for attempt in range(5):
        try:
            r = requests.post(API, headers={"Authorization": "Bearer " + key,
                              "Content-Type": "application/json"}, json=payload, timeout=180)
            if r.status_code == 200:
                return r.json()["choices"][0]["message"]["content"]
            last = "http %s" % r.status_code
            if r.status_code not in (429, 500, 502, 503, 504):
                break
        except Exception as e:
            last = str(e)[:80]
        time.sleep(2 * (attempt + 1))
    raise RuntimeError(last or "call failed")


def parse_obj(content):
    obj = None
    for cand in re.findall(r"\{[^{}]*\}", content, re.S):
        try:
            obj = json.loads(cand)
        except ValueError:
            continue
    if obj is None:
        raise ValueError("no json")
    return obj


def judge_opinion(text, row, key):
    rc = col(row, "relief")
    pc = col(row, "penalty")
    ic = col(row, "issue")
    dc = col(row, "dispositive", "sentence", "holding", "quote")
    jc = col(row, "judge")
    yc = col(row, "year")
    ac = col(row, "amount")
    user = (
        "OPINION TEXT:\n" + text + "\n\nRECORDED VALUES:\n"
        + "taxpayer_relief: " + str(row.get(rc, "") or "(none recorded)") + "\n"
        + "penalty_outcome: " + str(row.get(pc, "") or "(none recorded)") + "\n"
        + "principal_issue: " + str(row.get(ic, "") or "(none recorded)") + "\n"
        + "dispositive_sentence: " + str(row.get(dc, "") or "(none recorded)") + "\n"
        + "authoring_judge: " + str(row.get(jc, "") or "(none recorded)") + "\n"
        + "tax_years: " + str(row.get(yc, "") or "(none recorded)") + "\n"
        + "amount_at_issue: " + str(row.get(ac, "") or "(none recorded)") + "\n\n"
        "Read the opinion and decide the truth, then compare the recorded values.\n"
        "true_relief: the correct taxpayer relief level. none = the Commissioner prevails on every contested "
        "issue the court decides; partial = each side wins at least one contested issue; substantial = the "
        "taxpayer wins the principal issue or the relief sought. A concession by the Commissioner is not a "
        "taxpayer win.\n"
        "true_penalty: the correct penalty outcome. no_penalties = none at issue; upheld = all sustained; "
        "rejected = none sustained; split = some sustained and some not.\n"
        "relief_correct: true if recorded taxpayer_relief equals true_relief.\n"
        "penalty_correct: true if recorded penalty_outcome equals true_penalty.\n"
        "dispositive_valid: true only if dispositive_sentence appears in the opinion, states the court's holding "
        "on the principal issue, and is not the boilerplate decision line such as 'Decision will be entered under "
        "Rule 155' or 'An appropriate order and decision will be entered'.\n"
        "issue_correct: true if principal_issue names the main issue the case turned on.\n"
        "judge_correct: true if authoring_judge names the judge who authored the opinion.\n"
        "years_correct: true if tax_years matches the tax year or years at issue in the opinion.\n"
        "amount_correct: true if amount_at_issue matches the total deficiency or amount in dispute stated in the "
        "opinion, or is reasonably 'not stated' when the opinion gives no such figure.\n"
        'Reply with only this JSON: {"true_relief": "none|partial|substantial", "relief_correct": <bool>, '
        '"true_penalty": "no_penalties|upheld|rejected|split", "penalty_correct": <bool>, '
        '"dispositive_valid": <bool>, "issue_correct": <bool>, "judge_correct": <bool>, '
        '"years_correct": <bool>, "amount_correct": <bool>}.')
    o = parse_obj(call_model(user, key))
    return {
        "true_relief": norm(o.get("true_relief")) or "none",
        "true_penalty": norm(o.get("true_penalty")) or "no_penalties",
        "relief_correct": bool(o.get("relief_correct")),
        "penalty_correct": bool(o.get("penalty_correct")),
        "dispositive_valid": bool(o.get("dispositive_valid")),
        "issue_correct": bool(o.get("issue_correct")),
        "judge_correct": bool(o.get("judge_correct")),
        "years_correct": bool(o.get("years_correct")),
        "amount_correct": bool(o.get("amount_correct")),
    }


def judge_deliverables(brief, summary, key):
    if len(brief.strip()) < 200 and summary == "(no summary sheet provided)":
        return 0.0
    user = (
        "A worker produced a short brief and a summary sheet for a set of United States Tax Court memorandum "
        "opinions. Judge them on their own terms.\n\nBRIEF:\n" + brief[:7000] + "\n\nSUMMARY SHEET ROWS:\n"
        + summary + "\n\n"
        "grounded: true if the brief discusses specific outcomes, issues, or penalties rather than generic filler.\n"
        "breakdown: true if the brief describes how the cases break down across the set.\n"
        "summary_ok: true if the summary sheet gives a sensible rollup such as the total count and counts per "
        "relief level or penalty outcome.\n"
        'Reply with only this JSON: {"grounded": <bool>, "breakdown": <bool>, "summary_ok": <bool>}.')
    o = parse_obj(call_model(user, key))
    return 0.45 * bool(o.get("grounded")) + 0.30 * bool(o.get("breakdown")) + 0.25 * bool(o.get("summary_ok"))


def balanced(verds, true_key, ok_key):
    counts = {}
    for v in verds:
        counts[v[true_key]] = counts.get(v[true_key], 0) + 1
    k = max(1, len(counts))
    return sum(1.0 / (k * counts[v[true_key]]) for v in verds if v[ok_key])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--agent-dir", default="/logs/agent")
    ap.add_argument("--reward-out", default="/logs/verifier/reward.txt")
    a = ap.parse_args()
    lines = []

    def finish(score):
        score = round(max(0.0, min(1.0, score)), 4)
        os.makedirs(os.path.dirname(a.reward_out) or ".", exist_ok=True)
        open(a.reward_out, "w").write(str(score))
        body = "\n".join(["Score: %s" % score, ""] + lines)
        for p in ("/logs/verifier/verify_justification.txt", os.path.join(a.agent_dir, "verify_justification.txt")):
            try:
                os.makedirs(os.path.dirname(p), exist_ok=True)
                open(p, "w").write(body)
            except OSError:
                pass
        print(body)

    key = os.environ.get("FIREWORKS_API_KEY") or os.environ.get("OPENAI_API_KEY")
    if not key:
        lines.append("No judge API key in environment; failing closed.")
        finish(0.0)
        return
    try:
        cohort = json.load(open(COHORT))
    except (OSError, ValueError):
        lines.append("No cohort.json.")
        finish(0.0)
        return
    rows = read_table(a.agent_dir)
    if not rows:
        lines.append("No usable opinions table (.xlsx or .csv) in agent dir.")
        finish(0.0)
        return

    id_col = col(rows[0], "opinion_id", "id")
    cap_col = col(rows[0], "caption", "case name", "case_name", "case")
    by_id, cap_rows = {}, []
    for r in rows:
        rid = norm(r.get(id_col, "")) if id_col else ""
        if rid:
            by_id.setdefault(rid, r)
        if cap_col:
            cap_rows.append((cap_tokens(r.get(cap_col, "")), r))

    def match(entry):
        if norm(entry.get("id", "")) in by_id:
            return by_id[norm(entry["id"])]
        gt = cap_tokens(entry.get("caption", ""))
        best, bs, bo = None, 0, 0.0
        for at, r in cap_rows:
            if not at or not gt:
                continue
            shared = len(at & gt)
            ov = shared / min(len(at), len(gt))
            if (shared, ov) > (bs, bo):
                bs, bo, best = shared, ov, r
        return best if (bs >= 2 or (bs >= 1 and bo >= 0.5)) else None

    def work(entry):
        text = corpus_text(entry["id"])
        if not text:
            return None
        row = match(entry) or {}
        try:
            return judge_opinion(text, row, key)
        except Exception:
            return "FAIL"

    out = list(ThreadPoolExecutor(max_workers=MAXWORKERS).map(work, cohort))
    verds = [v for v in out if isinstance(v, dict)]
    failed = sum(1 for v in out if v == "FAIL")
    if failed > max(2, len(cohort) // 10):
        lines.append("Judge unavailable for %d/%d opinions; failing closed." % (failed, len(cohort)))
        finish(0.0)
        return

    n = len(cohort)
    relief = balanced(verds, "true_relief", "relief_correct")
    penalty = balanced(verds, "true_penalty", "penalty_correct")
    disp = sum(1 for v in verds if v["dispositive_valid"]) / n
    issue = sum(1 for v in verds if v["issue_correct"]) / n
    judge = sum(1 for v in verds if v["judge_correct"]) / n
    years = sum(1 for v in verds if v["years_correct"]) / n
    amount = sum(1 for v in verds if v["amount_correct"]) / n
    op = (W["relief"] * relief + W["penalty"] * penalty + W["dispositive"] * disp + W["issue"] * issue
          + W["judge"] * judge + W["years"] * years + W["amount"] * amount)

    deliv_error = ""
    try:
        deliv = judge_deliverables(read_brief(a.agent_dir), read_summary(a.agent_dir), key)
    except Exception as e:
        deliv = 0.0
        deliv_error = str(e)[:160]

    reward = OPINIONS_WEIGHT * op + DELIVERABLES_WEIGHT * deliv
    lines += [
        "opinions %.4f x %.2f + deliverables %.4f x %.2f" % (op, OPINIONS_WEIGHT, deliv, DELIVERABLES_WEIGHT),
        "balanced relief %.3f penalty %.3f | dispositive %.3f issue %.3f judge %.3f years %.3f amount %.3f"
        % (relief, penalty, disp, issue, judge, years, amount),
        "cohort: %d; judged: %d; judge failures: %d" % (n, len(verds), failed),
    ]
    if deliv_error:
        lines.append("ERROR deliverables judge call failed (infrastructure, scored 0.0): " + deliv_error)
    finish(reward)


if __name__ == "__main__":
    main()

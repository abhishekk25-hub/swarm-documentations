#!/usr/bin/env python3


import glob
import json
import os
import re
import sys

from pptx import Presentation
from pptx.util import Pt

MAX_SLIDES_PER_CHAPTER = 5


def chap_num(path):
    m = re.search(r"chapter_(\d+)_slides", os.path.basename(path))
    return int(m.group(1)) if m else 9999


def coerce_slides(obj):
    """Return (chapter:int|None, title:str, slides:list[{title,bullets}])."""
    chapter = None
    title = ""
    slides = []
    if isinstance(obj, dict):
        chapter = obj.get("chapter")
        title = str(obj.get("title", "") or "")
        raw = obj.get("slides", [])
    elif isinstance(obj, list):
        raw = obj
    else:
        raw = []
    for s in raw if isinstance(raw, list) else []:
        if isinstance(s, dict):
            st = str(s.get("title", "") or "")
            b = s.get("bullets", s.get("points", []))
            if isinstance(b, str):
                bullets = [b]
            elif isinstance(b, list):
                bullets = [str(x) for x in b if str(x).strip()]
            else:
                bullets = []
            slides.append({"title": st, "bullets": bullets})
        elif isinstance(s, str):
            slides.append({"title": s, "bullets": []})
    return chapter, title, slides


def add_slide(prs, title, bullets):
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    # title
    if slide.shapes.title is not None:
        slide.shapes.title.text = title or ""
    else:
        tb = slide.shapes.add_textbox(Pt(40), Pt(30), Pt(640), Pt(60))
        tb.text_frame.text = title or ""
    # body bullets
    body = None
    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.idx != 0:  # not the title
                body = ph
                break
        except Exception:
            continue
    if body is None:
        body = slide.shapes.add_textbox(Pt(40), Pt(110), Pt(640), Pt(360))
    tf = body.text_frame
    tf.word_wrap = True
    if bullets:
        tf.text = bullets[0]
        for line in bullets[1:]:
            p = tf.add_paragraph()
            p.text = line
    return slide


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else "/logs/agent/output.pptx"
    src_dir = sys.argv[2] if len(sys.argv) > 2 else "/logs/agent"

    files = sorted(glob.glob(os.path.join(src_dir, "chapter_*_slides.json")), key=chap_num)
    if not files:
        raise SystemExit(f"assemble_deck: no chapter_NN_slides.json files found in {src_dir}")

    prs = Presentation()
    total = 0
    for fp in files:
        n = chap_num(fp)
        try:
            obj = json.loads(open(fp, encoding="utf-8").read())
        except Exception as e:
            print(f"assemble_deck: skip {fp}: {e}")
            continue
        chapter, title, slides = coerce_slides(obj)
        if chapter is None:
            chapter = n
        slides = slides[:MAX_SLIDES_PER_CHAPTER]
        if not slides:
            slides = [{"title": f"Chapter {chapter}: {title}".rstrip(": "), "bullets": []}]
        # force the first slide of the chapter to carry the chapter marker
        head = f"Chapter {chapter}: {title}".rstrip(": ").strip()
        first_title = slides[0]["title"].strip()
        if not re.match(rf"^\s*chapter\s*[#:\-]?\s*0*{chapter}\b", first_title, re.IGNORECASE):
            slides[0]["title"] = head if head else f"Chapter {chapter}"
        for s in slides:
            add_slide(prs, s["title"], s["bullets"])
            total += 1

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    prs.save(out_path)
    print(f"assemble_deck: stitched {len(files)} chapters / {total} slides -> {out_path}")


if __name__ == "__main__":
    main()

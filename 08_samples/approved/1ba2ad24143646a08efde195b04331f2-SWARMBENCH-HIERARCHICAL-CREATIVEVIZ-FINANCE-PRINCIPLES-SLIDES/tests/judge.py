#!/usr/bin/env python3

import argparse
import json
import os
import re
import subprocess
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed

MODEL = "accounts/fireworks/models/kimi-k2p5"
N_CHAPTERS = 20

# Tunables (overridable via env)
BATCH_CHAPTERS = int(os.environ.get("JUDGE_BATCH", "3"))
MAX_WORKERS = int(os.environ.get("JUDGE_WORKERS", "8"))
CHAP_CHARS = int(os.environ.get("JUDGE_CHAP_CHARS", "48000"))
SLIDES_CHAR_CAP = int(os.environ.get("JUDGE_SLIDES_CAP", "12000"))
PAD_PAGES = int(os.environ.get("JUDGE_PAD_PAGES", "4"))
FRONT_SKIP = int(os.environ.get("JUDGE_FRONT_SKIP", "16"))
LLM_RETRIES = 3

# Legacy fallback only
CHUNK_CHARS = 400000
CHUNK_OVERLAP = 60000

CHAPTER_RE = re.compile(r"^\s*chapter\s*[#:\-]?\s*0*([0-9]{1,2})\b", re.IGNORECASE)
FIG_RE = re.compile(r"\b(?:FIGURE|TABLE)\s+(\d{1,2})\.\d", re.IGNORECASE)
BACKMATTER_RE = re.compile(r"(?im)^\s*(answer key|references|^index\b|appendix\s+[a-z]\b)")


def write_reward(path, reward, justification):
    reward = max(0.0, min(1.0, float(reward)))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"reward": reward}, f)
    for jpath in ("/logs/agent/judge_justification.txt", "/logs/verifier/judge_justification.txt"):
        try:
            os.makedirs(os.path.dirname(jpath), exist_ok=True)
            with open(jpath, "w", encoding="utf-8") as f:
                f.write(f"Score: {reward}\n\n{justification}\n")
        except Exception:
            pass


def extract_json(text):
    text = (text or "").strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", text)
    if m:
        text = m.group(1).strip()
    m = re.search(r"\{[\s\S]*\}", text)
    if m:
        return m.group(0)
    return text


def _slide_text(slide):
    """Return (title, full_text) for one slide, gathering every text frame."""
    title = ""
    try:
        if slide.shapes.title is not None and (slide.shapes.title.text or "").strip():
            title = slide.shapes.title.text.strip()
    except Exception:
        title = ""

    lines = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for para in shape.text_frame.paragraphs:
            t = "".join(run.text for run in para.runs).strip()
            if not t:
                t = (para.text or "").strip()
            if t:
                lines.append(t)

    if not title and lines:
        title = lines[0]
    body = "\n".join(lines)
    return title, body


def load_agent_slides(path):
    """Read the .pptx and return {chapter_number: combined_slide_text}."""
    from pptx import Presentation
    prs = Presentation(path)

    sections = {}
    current = None
    buf = []

    def flush():
        if current is not None and buf:
            sections.setdefault(current, "")
            sections[current] += ("\n".join(buf) + "\n")

    for slide in prs.slides:
        title, body = _slide_text(slide)
        m = CHAPTER_RE.match(title or "")
        if m:
            ch = int(m.group(1))
            if 1 <= ch <= N_CHAPTERS:
                flush()
                current = ch
                buf = [body if body else title]
                continue
        if current is not None:
            if body:
                buf.append(body)
    flush()

    sections = {c: t.strip() for c, t in sections.items() if t.strip()}
    if not sections:
        raise ValueError("no 'Chapter N' titled slides were found in the .pptx")
    return sections


def extract_book_pages(pdf_path):
    """Return the book as a list of per-page text strings (pdftotext, pypdf fallback)."""
    try:
        out = subprocess.run(["pdftotext", "-layout", pdf_path, "-"],
                             capture_output=True, timeout=600)
        if out.returncode == 0 and out.stdout:
            text = out.stdout.decode("utf-8", errors="ignore")
            pages = text.split("\x0c")
            if len(pages) > 50:
                return pages
    except Exception:
        pass
    try:
        from pypdf import PdfReader
        reader = PdfReader(pdf_path)
        return [(pg.extract_text() or "") for pg in reader.pages]
    except Exception:
        return []


def detect_chapter_spans(pages):
    """Map each chapter 1..N to its text using FIGURE/TABLE 'k.M' anchors."""
    if not pages or len(pages) < 50:
        return None

    anchor = {}
    for i, t in enumerate(pages):
        if i < FRONT_SKIP:
            continue
        for m in FIG_RE.finditer(t):
            k = int(m.group(1))
            if 1 <= k <= N_CHAPTERS and k not in anchor:
                anchor[k] = i

    if len(anchor) < N_CHAPTERS - 2:
        return None
    ordered = [anchor.get(k) for k in range(1, N_CHAPTERS + 1)]
    seen = [p for p in ordered if p is not None]
    if any(seen[i] >= seen[i + 1] for i in range(len(seen) - 1)):
        return None

    for k in range(1, N_CHAPTERS + 1):
        if anchor.get(k) is None:
            prev = next((anchor[j] for j in range(k - 1, 0, -1) if anchor.get(j) is not None), FRONT_SKIP)
            nxt = next((anchor[j] for j in range(k + 1, N_CHAPTERS + 1) if anchor.get(j) is not None), len(pages))
            anchor[k] = (prev + nxt) // 2

    gaps = [anchor[k + 1] - anchor[k] for k in range(1, N_CHAPTERS)]
    med = sorted(gaps)[len(gaps) // 2] if gaps else 24

    starts = {}
    prev = FRONT_SKIP - 1
    for k in range(1, N_CHAPTERS + 1):
        s = max(prev + 1, anchor[k] - PAD_PAGES, FRONT_SKIP)
        starts[k] = s
        prev = s

    back = None
    for i in range(anchor[N_CHAPTERS] + 5, len(pages)):
        if BACKMATTER_RE.search(pages[i]):
            back = i
            break
    end_last = back if back is not None else min(len(pages), anchor[N_CHAPTERS] + med)

    spans = {}
    for k in range(1, N_CHAPTERS + 1):
        end = starts[k + 1] if k < N_CHAPTERS else end_last
        text = "\n".join(pages[starts[k]:end]).strip()
        if text:
            spans[k] = text
    return spans if len(spans) >= N_CHAPTERS - 2 else None


def _client():
    from openai import OpenAI
    return OpenAI(api_key=os.environ["FIREWORKS_API_KEY"],
                  base_url="https://api.fireworks.ai/inference/v1")


def _llm_json(client, prompt):
    last_err = None
    for _ in range(LLM_RETRIES):
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
            )
            raw = resp.choices[0].message.content or ""
            return json.loads(extract_json(raw))
        except Exception as e:
            last_err = e
    raise RuntimeError(f"{type(last_err).__name__}: {last_err}")


RUBRIC = """You are a demanding grader scoring a student's presentation slides, each set being the
key-point slides for one chapter of a Principles of Finance textbook. Grade like a tough
professor, not a generous one.

For each chapter below you are given (A) the student's slides for that chapter (a small deck
of at most 5 slides: titles plus bullet points) and (B) the actual textbook text for that
same chapter. Grade each chapter independently.

Remember this is a SLIDE DECK, not an essay: concise, well-chosen bullet points are expected
and good. Do NOT penalize brevity itself. Judge whether the slides capture the chapter's real
substance.

Method (do this internally per chapter, then output only json):
  1. From the chapter's book text, enumerate that chapter's major key points, the specific
     concepts, named models, key definitions, formulas, and important distinctions that a
     faithful key-point deck of this chapter must include (most chapters have several, often
     6-10 such points).
  2. Read the student's slides and check how many of those major points the slides actually
     capture with specificity and correctness, using the chapter's real terminology, models,
     and formulas, not vague bullets that could apply to any chapter and not mere topic labels
     with no substance.
  3. Separately check faithfulness: any bullet that is fabricated, wrong, or unsupported by the
     chapter text is a serious defect.

Scoring scale:
  - 0.90-1.00: The slides capture essentially ALL of the chapter's major key points with
    precise, correct, chapter-specific content; nothing important missing and nothing
    fabricated. A presenter could teach the whole chapter from them.
  - 0.70-0.89: Captures most major points but misses at least one important concept, OR several
    bullets are generic/surface-level rather than substantive.
  - 0.50-0.69: Conveys the broad gist but misses multiple major points, stays mostly high-level
    or generic, or merely lists topics without the chapter's specific terms/models/formulas.
  - 0.30-0.49: Only superficial or partial coverage; several major concepts absent.
  - 0.01-0.29: Minimal, mostly generic, nearly empty, or largely off-target.
  - 0.00: No usable slides for the chapter, or content that is wrong/fabricated/not from the chapter.

Judge solely against the provided chapter text, not your own outside knowledge.

Respond with only a JSON object:
{"chapter_scores": {"<chapter number>": {"score": <float 0..1>, "missing": ["<major point omitted>", "..."], "note": "<one short sentence justifying the score>"}}}

"""


def _parse_scores(parsed):
    scores = parsed.get("chapter_scores", parsed)
    out = {}
    if not isinstance(scores, dict):
        return out
    for k, v in scores.items():
        try:
            ch = int(re.sub(r"[^0-9]", "", str(k)))
        except Exception:
            continue
        if not (1 <= ch <= N_CHAPTERS):
            continue
        if isinstance(v, dict):
            sc = float(v.get("score", 0.0))
            note = str(v.get("note", ""))[:200]
            missing = v.get("missing")
            if isinstance(missing, list) and missing:
                note = (note + " | missing: " + "; ".join(str(x) for x in missing))[:400]
        else:
            sc = float(v)
            note = ""
        out[ch] = (max(0.0, min(1.0, sc)), note)
    return out


def score_batch(client, batch, sections, spans):
    parts = []
    for c in batch:
        parts.append(
            f"==== CHAPTER {c} ====\n"
            f", (A) STUDENT SLIDES (Chapter {c}) ---\n{sections[c][:SLIDES_CHAR_CAP]}\n\n"
            f", (B) BOOK TEXT (Chapter {c}) ---\n{spans[c][:CHAP_CHARS]}"
        )
    prompt = RUBRIC + "\n\n".join(parts)
    return _parse_scores(_llm_json(client, prompt))


def chapter_sliced_judge(sections, spans):
    client = _client()
    gradable = sorted(c for c in range(1, N_CHAPTERS + 1) if c in sections and c in spans)
    batches = [gradable[i:i + BATCH_CHAPTERS] for i in range(0, len(gradable), BATCH_CHAPTERS)]

    best = {}
    errors = 0
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = {ex.submit(score_batch, client, b, sections, spans): b for b in batches}
        for fut in as_completed(futs):
            try:
                res = fut.result()
            except Exception:
                errors += 1
                continue
            for ch, (sc, note) in res.items():
                if ch not in best or sc > best[ch][0]:
                    best[ch] = (sc, note)

    if not best:
        return None, errors

    reward = sum(best.get(c, (0.0, ""))[0] for c in range(1, N_CHAPTERS + 1)) / N_CHAPTERS
    breakdown = {str(c): {"score": round(best[c][0], 4), "note": best[c][1]} for c in sorted(best)}
    justification = (
        "Chapter-sliced + batched + parallel LLM judging of the .pptx slide deck. The book was "
        "split into per-chapter text blocks via the textbook's FIGURE/TABLE 'k.M' numbering; each "
        "chapter's slides were graded against ONLY that chapter's own text. "
        f"{len(batches)} batches of up to {BATCH_CHAPTERS} chapters ran concurrently "
        f"({MAX_WORKERS} workers). Final reward = mean of per-chapter scores over {N_CHAPTERS} "
        f"chapters (unscored = 0). Chapters scored: {len(best)}/{N_CHAPTERS}; batch errors: "
        f"{errors}.\n\n" + json.dumps(breakdown, indent=2)
    )
    return (reward, justification), errors


MAP_PROMPT = """Below is a contiguous excerpt from a Principles of Finance textbook whose chapters are
numbered 1 to 20. Identify which chapters have their FULL section present in this excerpt
(beginning through end-of-chapter material). Do not include a chapter cut off at the start or end.
Respond with only json: {"chapters_fully_present": [<chapter numbers>]}

EXCERPT:
"""


def _legacy_chunks(text):
    chunks = []
    step = max(1, CHUNK_CHARS - CHUNK_OVERLAP)
    i = 0
    while i < len(text):
        chunks.append(text[i:i + CHUNK_CHARS])
        if i + CHUNK_CHARS >= len(text):
            break
        i += step
    return chunks


def legacy_judge(sections, book_text):
    client = _client()

    def map_chunk(chunk):
        parsed = _llm_json(client, MAP_PROMPT + chunk)
        vals = parsed.get("chapters_fully_present", parsed if isinstance(parsed, list) else [])
        present = set()
        if isinstance(vals, list):
            for x in vals:
                try:
                    c = int(re.sub(r"[^0-9]", "", str(x)))
                    if 1 <= c <= N_CHAPTERS:
                        present.add(c)
                except Exception:
                    continue
        return present

    def score_chunk(chunk, present):
        relevant = {c: sections[c] for c in present if c in sections}
        if not relevant:
            return {}
        sp = {c: chunk for c in relevant}
        return score_batch(client, sorted(relevant), relevant, sp)

    chunks = _legacy_chunks(book_text)
    best = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        present_map = list(ex.map(lambda c: (c, map_chunk(c)), chunks))
        futs = {ex.submit(score_chunk, c, pres): c for c, pres in present_map}
        for fut in as_completed(futs):
            try:
                res = fut.result()
            except Exception:
                continue
            for ch, (sc, note) in res.items():
                if ch not in best or sc > best[ch][0]:
                    best[ch] = (sc, note)

    if not best:
        return None
    reward = sum(best.get(c, (0.0, ""))[0] for c in range(1, N_CHAPTERS + 1)) / N_CHAPTERS
    breakdown = {str(c): {"score": round(best[c][0], 4), "note": best[c][1]} for c in sorted(best)}
    just = ("LEGACY chunk+map+score judging (chapter detection was unreliable). Reward = mean "
            f"over {N_CHAPTERS} chapters (unscored = 0). Chapters scored: {len(best)}/{N_CHAPTERS}.\n\n"
            + json.dumps(breakdown, indent=2))
    return reward, just


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent-output", default="/logs/agent/output.pptx")
    parser.add_argument("--reward-out", default="/logs/verifier/reward.json")
    parser.add_argument("--pdf", default="/input_artifacts/Principles_Finance.pdf")
    args = parser.parse_args()

    try:
        sections = load_agent_slides(args.agent_output)
    except Exception as e:
        write_reward(args.reward_out, 0.0,
                     f"Agent .pptx missing/unreadable or has no 'Chapter N' titled slides "
                     f"(fail-closed): {type(e).__name__}: {e}")
        return

    try:
        import openai  # noqa: F401
    except Exception as e:
        write_reward(args.reward_out, 0.0,
                     f"LLM judge dependency unavailable (fail-closed): {type(e).__name__}: {e}")
        return

    if not os.environ.get("FIREWORKS_API_KEY"):
        write_reward(args.reward_out, 0.0,
                     "LLM judge could not run because FIREWORKS_API_KEY is not set (fail-closed).")
        return

    pages = extract_book_pages(args.pdf)
    if not pages or sum(len(p) for p in pages) < 1000:
        write_reward(args.reward_out, 0.0,
                     f"Could not extract book text from {args.pdf} (fail-closed).")
        return

    try:
        spans = detect_chapter_spans(pages)
        if spans:
            result, _errors = chapter_sliced_judge(sections, spans)
            if result is not None:
                write_reward(args.reward_out, result[0], result[1])
                return
        legacy = legacy_judge(sections, "\n".join(pages))
        if legacy is None:
            write_reward(args.reward_out, 0.0,
                         "LLM judge returned no usable chapter scores (fail-closed).")
            return
        write_reward(args.reward_out, legacy[0], legacy[1])
    except Exception as e:
        tb = traceback.format_exc(limit=8)
        write_reward(args.reward_out, 0.0,
                     f"LLM judge error (fail-closed): {type(e).__name__}: {e}\n{tb}")
        return


if __name__ == "__main__":
    main()

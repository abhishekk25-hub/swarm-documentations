import csv
import json
import os
import re

MODEL = "accounts/fireworks/models/kimi-k2p5"
AGENT_DIR = os.environ.get("AGENT_LOG_DIR", "/logs/agent")
ORACLE_PATH = os.environ.get("ORACLE_PATH", "/tests/oracle.json")

PARTS = ["status_high_risk", "recruiting_stale", "completed_no_results", "overdue_primary_completion"]
PART_KW = ("score", "part", "status", "stale", "recruit", "complet", "result", "overdue", "component")
SCHEDULE_COLS = ["site_id", "partner_org", "region", "modality", "proposed_month",
                 "travel_nights", "risk_score", "objective_contribution", "selection_rank"]
CONSTRAINT_NAMES = ["visit", "travel night", "partner coverage", "one visit per site", "roster coverage"]


def truthy(v):
    return str(v).strip().lower() in {"true", "yes", "1", "y", "t", "pass", "passed", "ok"}


def as_int(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def read_csv(path):
    if not os.path.isfile(path):
        return [], []
    with open(path, newline="", encoding="utf-8") as f:
        r = csv.DictReader(f)
        return (r.fieldnames or []), list(r)


def has_col(cols, *frags):
    return any(any(f in c.lower() for f in frags) for c in cols)


def find_col(cols, *frags, exclude=()):
    for c in cols:
        lc = c.lower()
        if any(f in lc for f in frags) and not any(e in lc for e in exclude):
            return c
    return None


def total_col(cols):
    # The combined risk score, not one of the parts.
    for c in cols:
        if c.lower() in ("risk_score", "total_risk_score", "risk", "total_risk"):
            return c
    return find_col(cols, "risk", exclude=("status", "stale", "result", "overdue", "part"))


def part_cols(cols, rows):
    total = total_col(cols)
    out = []
    for c in cols:
        lc = c.lower()
        if c == total or "total" in lc or not any(k in lc for k in PART_KW):
            continue
        vals = [as_int(r.get(c)) for r in rows[:10] if (r.get(c) or "").strip() != ""]
        if vals and all(v is not None and 0 <= v <= 3 for v in vals):
            out.append(c)
    return out


def filled_fraction(rows, cols, *frags):
    col = next((c for c in cols if any(f in c.lower() for f in frags)), None)
    if not col or not rows:
        return 0.0
    return sum(1 for r in rows if (r.get(col) or "").strip()) / len(rows)


def load_oracle():
    with open(ORACLE_PATH, encoding="utf-8") as f:
        return json.load(f)

def score_static(oracle):
    sites = {s["site_id"]: s for s in oracle["sites"]}
    cons = oracle["constraints"]
    partners = {s["partner_org"] for s in sites.values()}
    opt = oracle["optimal_plan"]["total_risk"]
    n = len(sites)

    a_cols, a_rows = read_csv(os.path.join(AGENT_DIR, "audit_trace.csv"))
    s_cols, s_rows = read_csv(os.path.join(AGENT_DIR, "monitoring_schedule.csv"))
    c_cols, c_rows = read_csv(os.path.join(AGENT_DIR, "constraint_checks.csv"))

    audit = {(r.get("site_id") or "").strip(): r for r in a_rows if (r.get("site_id") or "").strip()}

    a_total = total_col(a_cols)
    a_part_cols = part_cols(a_cols, a_rows)

    risk_ok = parts_ok = 0
    for sid, s in sites.items():
        r = audit.get(sid)
        if not r:
            continue
        reported = as_int(r.get(a_total)) if a_total else None
        if reported == s["risk_score"]:
            risk_ok += 1
        vals = [as_int(r.get(c)) for c in a_part_cols]
        if len(a_part_cols) >= 4 and None not in vals and sum(vals) == reported:
            parts_ok += 1
    risk_accuracy = risk_ok / n if n else 0.0
    parts_separated = parts_ok / n if n else 0.0

    sched_ids = [(r.get("site_id") or "").strip() for r in s_rows]
    real_ids = [sid for sid in sched_ids if sid in sites]
    selected = set(real_ids)
    one_per_site = bool(sched_ids) and len(sched_ids) == len(set(sched_ids)) and len(real_ids) == len(sched_ids)
    visits = len(selected)
    nights = sum(sites[sid]["travel_nights"] for sid in selected)
    covered = {sites[sid]["partner_org"] for sid in selected}
    true_risk = sum(sites[sid]["risk_score"] for sid in selected)

    within_visits = visits <= cons["max_visits"]
    within_nights = nights <= cons["max_nights"]
    all_covered = covered == partners
    plan_valid = bool(selected) and one_per_site and within_visits and within_nights and all_covered

    # Graduated plan quality: start from the share of optimal risk the plan captures, then
    # apply a soft penalty for each rule it bends so a near-miss (e.g. 19 visits instead of
    # 18, or one partner short) still earns most of the credit instead of collapsing to zero.
    coverage = min(1.0, true_risk / opt) if opt else 0.0
    pen_visits = 1.0 if within_visits else max(0.0, 1.0 - (visits - cons["max_visits"]) / cons["max_visits"])
    pen_nights = 1.0 if within_nights else max(0.0, 1.0 - (nights - cons["max_nights"]) / cons["max_nights"])
    pen_partner = len(covered) / len(partners) if partners else 0.0
    pen_unique = len(set(real_ids)) / len(sched_ids) if sched_ids else 0.0
    plan_quality = coverage * pen_visits * pen_nights * pen_partner * pen_unique if selected else 0.0
    plan_quality = max(0.0, min(1.0, plan_quality))

    s_modality = find_col(s_cols, "modality", "study_type", "study type")
    mod_ok = sum(1 for r in s_rows
                 if (r.get("site_id") or "").strip() in sites and s_modality
                 and (r.get(s_modality) or "").strip().lower() == sites[(r.get("site_id") or "").strip()]["modality"].lower())
    modality_accuracy = mod_ok / len(real_ids) if real_ids else 0.0

    a_sel = find_col(a_cols, "select", exclude=("rank",))
    flagged = {sid for sid, r in audit.items() if a_sel and truthy(r.get(a_sel))}

    audit_completeness = sum([
        len(audit) == n,
        has_col(a_cols, "fda") and filled_fraction(a_rows, a_cols, "fda") >= 0.9,
        has_col(a_cols, "oig") and filled_fraction(a_rows, a_cols, "oig") >= 0.9,
        filled_fraction(a_rows, a_cols, "url", "source") >= 0.9,
        parts_separated >= 0.95,
        bool(a_sel) and flagged == selected,
    ]) / 6.0

    ranks = [as_int(r.get("selection_rank")) for r in s_rows]
    schedule_ok = sum([
        set(SCHEDULE_COLS).issubset(set(s_cols)) and len(part_cols(s_cols, s_rows)) >= 4,
        one_per_site,
        sorted(r for r in ranks if r is not None) == list(range(1, len(s_rows) + 1)),
        filled_fraction(s_rows, s_cols, "objective_contribution") >= 0.9,
        modality_accuracy >= 0.95,
    ]) / 5.0

    truth = {
        "visit": within_visits,
        "travel night": within_nights,
        "partner coverage": all_covered,
        "one visit per site": one_per_site,
        "roster coverage": len(audit) == n,
    }
    def norm(s):
        return str(s).lower().replace("_", " ").replace("-", " ")

    constraint_score = 0.0
    for name in CONSTRAINT_NAMES:
        row = next((r for r in c_rows if name in norm(" ".join(map(str, r.values())))), None)
        if row and truthy(row.get("passed") or list(row.values())[-1]) == truth[name]:
            constraint_score += 1
    constraint_score /= len(CONSTRAINT_NAMES)

    static = (0.35 * risk_accuracy + 0.35 * plan_quality + 0.10 * audit_completeness +
              0.10 * schedule_ok + 0.10 * constraint_score)

    detail = {
        "risk_accuracy": round(risk_accuracy, 3),
        "parts_separated": round(parts_separated, 3),
        "modality_accuracy": round(modality_accuracy, 3),
        "plan_valid": plan_valid,
        "visits": visits, "nights": nights,
        "partners_covered": f"{len(covered)}/{len(partners)}",
        "selected_risk": true_risk, "optimal_risk": opt,
        "plan_quality": round(plan_quality, 3),
        "plan_penalties": {"visits": round(pen_visits, 3), "nights": round(pen_nights, 3),
                           "partner": round(pen_partner, 3), "unique": round(pen_unique, 3)},
        "audit_completeness": round(audit_completeness, 3),
        "schedule_ok": round(schedule_ok, 3),
        "constraint_score": round(constraint_score, 3),
    }
    return round(static, 4), detail


PROMPT = """You are an operations director reviewing a monitoring-visit memo one of your analysts wrote.
The memo defends this quarter's plan for which clinical-trial sites we send monitors to. Our rules
are simple: at most 18 visits and 30 travel nights, every partner organization gets at least one
visit, and beyond that we go after the riskiest sites first (risk comes from each trial's public
ClinicalTrials.gov record). Read it the way a busy director would and be honest, not generous.

Give each of these a score from 0 to 1:

coverage - does the memo actually explain how the schedule was chosen, which sites got dropped
because of the limits, how it broke ties, and the tradeoffs it made to keep every partner covered?

specificity - is it grounded in real detail (named sites or partners, the actual visit and night
totals, concrete risk reasoning) instead of vague filler?

consistency - do the numbers hold together and stay inside the caps (<=18 visits, <=30 nights), and
does it confirm every partner got a visit? If it describes a plan that breaks a cap or skips a
partner, this is low.

clarity - could I hand this to the board as-is? Is it a clean, well-organized memo?

Reply with json only:
{"coverage": <0-1>, "specificity": <0-1>, "consistency": <0-1>, "clarity": <0-1>, "notes": "<a sentence or two>"}
"""


def read_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if line:
                parts.append(line)
    return "\n".join(parts)


def extract_json(text):
    text = (text or "").strip()
    fence = re.search(r"```(?:json)?\s*([\s\S]*?)```", text)
    if fence:
        text = fence.group(1).strip()
    obj = re.search(r"\{[\s\S]*\}", text)
    return obj.group(0) if obj else text


def score_memo():
    path = os.path.join(AGENT_DIR, "scheduling_rationale.docx")
    if not os.path.isfile(path) or os.path.getsize(path) < 1000:
        return 0.0, {"error": "memo missing or too small"}
    try:
        text = read_docx(path)
    except Exception as e:
        return 0.0, {"error": f"could not read memo: {e}"}
    key = os.environ.get("FIREWORKS_API_KEY")
    if not key:
        return 0.0, {"error": "FIREWORKS_API_KEY not set"}
    try:
        from openai import OpenAI
        client = OpenAI(api_key=key, base_url="https://api.fireworks.ai/inference/v1")
    except Exception as e:
        return 0.0, {"error": f"openai client setup failed: {e}"}

    last = None
    for _ in range(3):
        try:
            resp = client.chat.completions.create(
                model=MODEL, temperature=0,
                messages=[{"role": "user", "content": PROMPT + "\n\nMEMO:\n" + text[:60000]}],
            )
            verdict = json.loads(extract_json(resp.choices[0].message.content or ""))
            dims = ["coverage", "specificity", "consistency", "clarity"]
            scores = {d: max(0.0, min(1.0, float(verdict.get(d, 0.0)))) for d in dims}
            return sum(scores.values()) / 4.0, {"scores": scores, "notes": verdict.get("notes", "")}
        except Exception as e:
            last = e
    return 0.0, {"error": f"llm call failed: {last}"}


def main():
    os.makedirs("/logs/verifier", exist_ok=True)
    try:
        oracle = load_oracle()
    except Exception as e:
        _write(0.0, {"error": f"oracle load failed: {e}"})
        return

    static, sdetail = score_static(oracle)
    memo, mdetail = score_memo()
    reward = round(0.6 * static + 0.4 * memo, 4)

    _write(reward, {
        "reward": reward,
        "static": {"score": static, "weight": 0.6, **sdetail},
        "memo": {"score": round(memo, 4), "weight": 0.4, **mdetail},
    })
    print(f"static {static}  memo {round(memo,3)}  ->  reward {reward}")
    print(json.dumps(sdetail, indent=2))


def _write(reward, payload):
    with open("/logs/verifier/reward.txt", "w") as f:
        f.write(str(reward))
    with open("/logs/verifier/reward.json", "w") as f:
        json.dump(payload, f, indent=2)
    try:
        with open("/logs/agent/judge_justification.txt", "w") as f:
            f.write(json.dumps(payload, indent=2))
    except OSError:
        pass


if __name__ == "__main__":
    main()

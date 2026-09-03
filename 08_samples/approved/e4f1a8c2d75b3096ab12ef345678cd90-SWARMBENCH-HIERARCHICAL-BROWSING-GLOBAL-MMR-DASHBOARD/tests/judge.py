import csv, io, json, os, re

MODEL = "accounts/fireworks/models/kimi-k2p5"
LLM_RETRIES = 3

BOARD_DIR = "/logs/agent/maternal-health-board"

AFRICA = [
    "NGA","SLE","TCD","SSD","COD","KEN","ETH","GHA","RWA","ZAF",
    "MLI","NER","AGO","MOZ","TZA","UGA","ZMB","ZWE","BFA","GIN",
    "CMR","SEN","CIV","MDG","MWI","SOM","SDN","ERI","LBR","TGO",
    "BEN","COG","CAF","GMB","MRT",
]
ASIA = [
    "IND","PAK","BGD","MMR","KHM","NPL","IDN","PHL","AFG","VNM",
    "LAO","TLS","PNG","LKA","THA","YEM","BTN","KGZ","TJK","UZB",
    "AZE","MNG","ARM","TKM","GEO",
]
AMERICAS = [
    "HTI","BOL","BRA","MEX","USA","PER","COL","VEN","GTM","HND",
    "NIC","SLV","PRY","ECU","DOM","CUB","JAM","TTO","GUY","SUR",
]
EUR_OCE = [
    "SWE","GBR","FRA","DEU","RUS","UKR","ROU","GRC","TUR","AUS",
    "POL","ITA","ESP","PRT","NLD","HUN","BGR","ALB","NZL","NOR",
]
ISO3_LIST = AFRICA + ASIA + AMERICAS + EUR_OCE
REGION_SLIDES = [AFRICA, ASIA, AMERICAS, EUR_OCE]
CHART_TITLES = ["global mmr comparison", "mmr vs health expenditure"]

REGION_COLORS = {
    "red":    lambda r, g, b: r > 150 and g < 100 and b < 100,
    "orange": lambda r, g, b: r > 180 and 80 < g < 180 and b < 80,
    "blue":   lambda r, g, b: r < 100 and g < 120 and b > 150,
    "green":  lambda r, g, b: r < 120 and g > 100 and b < 120,
}

def run_binary_checks():
    checks = []
    labels = []

    def chk(result, label):
        checks.append(bool(result))
        labels.append(label)

    def find_col(fields, *frags):
        for h in (fields or []):
            if any(f in h.lower() for f in frags):
                return h
        return None

    csv_path = os.path.join(BOARD_DIR, "data.csv")
    rows, fields = [], []
    if os.path.isfile(csv_path):
        with open(csv_path, newline="", encoding="utf-8") as f:
            rdr = csv.DictReader(f)
            fields = rdr.fieldnames or []
            rows = list(rdr)

    iso3_col = find_col(fields, "iso3")
    mmr_col  = next((h for h in fields
                     if "mmr" in h.lower() and "year" not in h.lower()
                     and "url" not in h.lower() and "source" not in h.lower()), None)

    chk(os.path.isfile(csv_path) and len(rows) == 100,
        "data.csv exists with 100 rows")
    chk(iso3_col is not None and mmr_col is not None,
        "data.csv has ISO3 and MMR columns")
    chk(find_col(fields, "who", "note") is not None
        and find_col(fields, "url", "source", "link") is not None,
        "data.csv has WHO note AND URL/source columns")

    present = {r[iso3_col].strip().upper() for r in rows if iso3_col and iso3_col in r}
    for region, label in [(AFRICA, "Africa"), (ASIA, "Asia"),
                          (AMERICAS, "Americas"), (EUR_OCE, "Europe/Oceania")]:
        pct = sum(1 for c in region if c in present) / len(region)
        chk(pct >= 0.90, f"{label} coverage ≥ 90 %")

    trend_path = os.path.join(BOARD_DIR, "mmr_trend.csv")
    t_ok = False
    if os.path.isfile(trend_path):
        with open(trend_path, newline="", encoding="utf-8") as f:
            tr = csv.DictReader(f)
            tf = tr.fieldnames or []
            trows = list(tr)
        t_ok = (find_col(tf, "year") is not None and
                find_col(tf, "iso3", "country") is not None and
                len(trows) >= 200)

    audit_path = os.path.join(BOARD_DIR, "source_audit.csv")
    a_ok = False
    if os.path.isfile(audit_path):
        with open(audit_path, newline="", encoding="utf-8") as f:
            ar = csv.DictReader(f)
            af = ar.fieldnames or []
            arows = list(ar)
        a_ok = (find_col(af, "url", "link", "source") is not None and
                len(arows) >= 100)

    chk(t_ok and a_ok, "mmr_trend.csv and source_audit.csv present and correct")

    pptx_path = os.path.join(BOARD_DIR, "dashboard.pptx")
    pptx_ok   = os.path.isfile(pptx_path) and os.path.getsize(pptx_path) > 50_000

    chk(pptx_ok, "dashboard.pptx exists (>50 KB)")

    if pptx_ok:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            prs    = Presentation(pptx_path)
            slides = list(prs.slides)

            chk(len(slides) == 6, "dashboard.pptx has exactly 6 slides")

            def _check_table(slide_idx, region, rname):
                if slide_idx >= len(slides):
                    return False
                tbl_shape = next((s for s in slides[slide_idx].shapes
                                  if s.has_table), None)
                if tbl_shape is None:
                    return False
                tbl = tbl_shape.table
                if len(tbl.columns) < 5:
                    return False
                iso3_in_tbl = []
                for ri in range(1, len(tbl.rows)):
                    try:
                        cell = tbl.cell(ri, 1).text.strip().upper()
                        if cell:
                            iso3_in_tbl.append(cell)
                    except Exception:
                        pass
                correct = [c for c in region if c in set(iso3_in_tbl)]
                ordered = [c for c in iso3_in_tbl if c in set(region)]
                return correct == ordered and len(correct) >= len(region) * 0.70

            chk(_check_table(0, AFRICA, "Africa") and _check_table(1, ASIA, "Asia"),
                "Africa and Asia tables ordered correctly (≥70 % countries)")
            chk(_check_table(2, AMERICAS, "Americas") and _check_table(3, EUR_OCE, "Europe"),
                "Americas and Europe tables ordered correctly (≥70 % countries)")

            for i, expected in enumerate(CHART_TITLES, start=4):
                if i < len(slides):
                    slide_text = " ".join(
                        sh.text_frame.text
                        for sh in slides[i].shapes if sh.has_text_frame
                    ).lower()
                    chk(expected in slide_text,
                        f'Slide {i+1} title contains "{expected}"')
                else:
                    chk(False, f"Slide {i+1} missing")

            def _detect_chart_colors(slide):
                found = {name: False for name in REGION_COLORS}
                A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
                for shape in slide.shapes:
                    if shape.has_chart:
                        try:
                            for series in shape.chart.series:
                                srgb = series._element.find(
                                    f".//{{{A_NS}}}solidFill/{{{A_NS}}}srgbClr"
                                )
                                if srgb is None:
                                    continue
                                hx = srgb.get("val", "")
                                if len(hx) != 6:
                                    continue
                                rv = int(hx[0:2], 16)
                                gv = int(hx[2:4], 16)
                                bv = int(hx[4:6], 16)
                                for name, test in REGION_COLORS.items():
                                    if test(rv, gv, bv):
                                        found[name] = True
                        except Exception:
                            pass
                if not any(found.values()):
                    for shape in slide.shapes:
                        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                            try:
                                from PIL import Image
                                img = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
                                pixels = list(img.getdata())
                                for name, test in REGION_COLORS.items():
                                    if sum(1 for rv, gv, bv in pixels if test(rv, gv, bv)) > 50:
                                        found[name] = True
                            except Exception:
                                pass
                return found

            for slide_idx, slide_label in [(4, "Slide 5 (bar)"), (5, "Slide 6 (scatter)")]:
                if slide_idx < len(slides):
                    color_found = _detect_chart_colors(slides[slide_idx])
                    for color_name in REGION_COLORS:
                        chk(color_found[color_name],
                            f"{slide_label} uses {color_name} for region colour")
                else:
                    for color_name in REGION_COLORS:
                        chk(False, f"{slide_label} missing — {color_name} colour check")

            bar_sort_ok = True
            if len(slides) > 4:
                for shape in slides[4].shapes:
                    if shape.has_chart:
                        try:
                            vals = []
                            for series in shape.chart.series:
                                for pt in series.values:
                                    v = getattr(pt, "value", None)
                                    if v is not None:
                                        vals.append(float(v))
                            if len(vals) >= 5:
                                bar_sort_ok = all(
                                    vals[i] >= vals[i + 1] - 0.01
                                    for i in range(len(vals) - 1)
                                )
                        except Exception:
                            pass
            chk(bar_sort_ok, "Slide 5 bar chart sorted highest-to-lowest MMR")

            iso3_label_ok = True
            if len(slides) > 5:
                for shape in slides[5].shapes:
                    if shape.has_chart:
                        try:
                            import lxml.etree as ET
                            xml_str = ET.tostring(shape.chart._element, encoding="unicode")
                            xml_upper = xml_str.upper()
                            hits = sum(1 for c in ISO3_LIST if c in xml_upper)
                            iso3_label_ok = hits >= 5
                        except Exception:
                            pass
            chk(iso3_label_ok, "Slide 6 scatter has ISO3 labels in chart data")

        except Exception:
            already_done = sum(1 for l in labels if l.startswith(
                ("dashboard.pptx has", "Africa and", "Americas and", "Slide")))
            for _ in range(15 - already_done):
                chk(False, "pptx parse error")
    else:
        for name in ["6 slides", "Africa+Asia tables", "Americas+Europe tables",
                     "slide 5 title", "slide 6 title",
                     "slide5 red", "slide5 orange", "slide5 blue", "slide5 green",
                     "slide6 red", "slide6 orange", "slide6 blue", "slide6 green",
                     "bar chart sort", "scatter ISO3 labels"]:
            chk(False, f"pptx absent — {name}")

    # summary.docx structural checks

    docx_path = os.path.join(BOARD_DIR, "summary.docx")
    docx_ok   = os.path.isfile(docx_path) and os.path.getsize(docx_path) > 5_000

    chk(docx_ok, "summary.docx exists (>5 KB)")

    if docx_ok:
        try:
            from docx import Document
            doc      = Document(docx_path)
            text     = " ".join(p.text for p in doc.paragraphs)
            all_cells = " ".join(
                cell.text for table in doc.tables
                for row in table.rows for cell in row.cells
            )
            iso3_hits = sum(1 for c in ISO3_LIST if c in (text + " " + all_cells).upper())
            chk(iso3_hits >= 15,
                f"summary.docx contains ≥15 ISO3 codes ({iso3_hits} found)")

            invest_rows = 0
            for table in doc.tables:
                if not table.rows:
                    continue
                header = " ".join(
                    table.cell(0, c).text for c in range(len(table.columns))
                ).lower()
                if "rank" in header or "priority score" in header:
                    invest_rows = len(table.rows) - 1
                    break
            if invest_rows == 0:
                invest_text = " ".join(
                    p.text for p in doc.paragraphs
                    if any(kw in p.text.lower() for kw in
                           ["investment", "candidate", "priority", "score"])
                )
                invest_rows = sum(1 for c in ISO3_LIST if c in invest_text.upper())
            chk(invest_rows == 15,
                f"summary.docx investment table has exactly 15 ranked candidates ({invest_rows} found)")

        except Exception:
            chk(False, "summary.docx ISO3 count check failed")
            chk(False, "summary.docx investment candidates check failed")
    else:
        chk(False, "summary.docx absent — ISO3 count check")
        chk(False, "summary.docx absent — investment candidates check")

    assert len(checks) == 27, f"Expected 27 binary checks, got {len(checks)}"
    return checks, labels

# LLM JUDGE — anti-fabrication rubric for summary.docx content─────────────────────────────────────────────────────────────────────────────

RUBRIC = """You are a strict evaluator checking whether an AI agent faithfully used live World Bank and WHO data (not invented numbers) to produce a maternal health executive summary for 100 countries.

The summary.docx must cover these 7 mandatory sections:
1. Average MMR (maternal mortality ratio) across all 100 countries — a specific number
2. Top 3 highest AND top 3 lowest burden countries by MMR — 6 named countries with MMR values
3. Average MMR per region (Africa, Asia, Americas, Europe & Oceania) with the biggest inter-regional gap
4. Countries with missing MMR data — explicit list or confirmed "none missing"
5. Countries where skilled birth attendance is below 50% — specific country names
6. Fifteen investment candidates ranked by priority score = MMR ÷ health expenditure (higher ratio = more impact per dollar)
7. MMR trend direction — improving vs worsening, with examples or counts

Score these 4 dimensions. Be strict: generic text, missing sections, or implausible numbers should score low.

SECTION_COMPLETENESS (0.0-1.0)
All 7 sections present with substantive, specific content (not generic placeholder text).
  1.0 = all 7 clearly addressed with real numbers and country names
  0.7 = 5-6 sections present
  0.4 = 3-4 sections present
  0.0 = fewer than 3 sections, or mostly filler

DATA_PLAUSIBILITY (0.0-1.0)
Numbers must be consistent with real-world maternal health data.
Real-world anchors to check against:
  - Sub-Saharan African countries (e.g. Chad, Nigeria, Sierra Leone) typically have MMR 300-1800
  - European/Oceania countries (e.g. Sweden, Norway, New Zealand) typically have MMR 2-15
  - Global average across a mix of 100 countries should be roughly 100-300
  - Top investment candidates (highest MMR/health_exp ratio) are almost always poor Sub-Saharan African or South Asian countries; wealthy countries never rank first
  Score 0.0 if: European countries listed with MMR > 100; African countries listed with MMR < 30; global average outside 50-500; all countries clustered at the same MMR value (strong fabrication signal); or investment table led by high-income countries.
  Score 1.0 if: all ranges are plausible and consistent with the above anchors.

INVESTMENT_VALIDITY (0.0-1.0)
The 15-candidate ranked table must be logically consistent with the priority formula (MMR ÷ health_expenditure).
  - Higher MMR + lower health spending -> higher priority score -> higher rank
  - Expected top candidates: Sub-Saharan African or South Asian countries (high burden, low spending)
  - Priority scores should decrease from rank 1 to rank 15
  Score 0.0 if: table absent; wealthy countries (Sweden, USA, Australia) ranked #1-5; scores increase rather than decrease; or fewer than 10 candidates listed.
  Score 1.0 if: 15 countries ranked, scores plausibly decrease, top candidates are high-MMR low-spending nations.

REGIONAL_COHERENCE (0.0-1.0)
Regional breakdown must be internally consistent.
  Expected ordering: Africa (highest MMR) > Asia > Americas > Europe/Oceania (lowest MMR)
  The biggest inter-regional gap must be between Africa and Europe/Oceania.
  Score 0.0 if: fewer than 3 regions mentioned; Europe has higher average MMR than Africa; or all regions report the same average.
  Score 1.0 if: all 4 regions present, ordering is correct, biggest gap is Africa-Europe.

Respond with JSON only — no other text:
{"section_completeness": <0-1>, "data_plausibility": <0-1>, "investment_validity": <0-1>, "regional_coherence": <0-1>, "notes": "<one or two sentences on what the document does well and what it fails at>"}
"""


def extract_json(text):
    text = (text or "").strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", text)
    if m:
        text = m.group(1).strip()
    m = re.search(r"\{[\s\S]*\}", text)
    if m:
        return m.group(0)
    return text


def read_docx_text(path):
    from docx import Document
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def call_llm(client, docx_text):
    prompt = RUBRIC + "\n\n=== DOCUMENT TEXT ===\n" + docx_text[:60000]
    last_err = None
    for _ in range(LLM_RETRIES):
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
            )
            raw = resp.choices[0].message.content or ""
            return json.loads(extract_json(raw))
        except Exception as e:
            last_err = e
    raise RuntimeError(f"LLM call failed after {LLM_RETRIES} retries: {last_err}")


def run_llm_judge():
    """Returns (llm_score, detail_dict). llm_score is 0.0 on any failure."""
    docx_path = os.path.join(BOARD_DIR, "summary.docx")

    if not os.path.isfile(docx_path) or os.path.getsize(docx_path) < 1000:
        return 0.0, {"error": "summary.docx missing or too small"}

    try:
        docx_text = read_docx_text(docx_path)
    except Exception as e:
        return 0.0, {"error": f"Could not read summary.docx: {e}"}

    api_key = os.environ.get("FIREWORKS_API_KEY")
    if not api_key:
        return 0.0, {"error": "FIREWORKS_API_KEY not set"}

    try:
        from openai import OpenAI
    except Exception as e:
        return 0.0, {"error": f"openai package unavailable: {e}"}

    client = OpenAI(api_key=api_key, base_url="https://api.fireworks.ai/inference/v1")

    try:
        result = call_llm(client, docx_text)
    except Exception as e:
        return 0.0, {"error": f"LLM call failed: {e}"}

    dims = ["section_completeness", "data_plausibility", "investment_validity", "regional_coherence"]
    scores = {d: max(0.0, min(1.0, float(result.get(d, 0.0)))) for d in dims}
    llm_score = sum(scores.values()) / len(dims)

    return llm_score, {"scores": scores, "notes": result.get("notes", "")}

def main():
    os.makedirs("/logs/verifier", exist_ok=True)

    checks, labels = run_binary_checks()
    binary_passed  = sum(checks)
    binary_total   = len(checks)  # 27

    llm_score, llm_detail = run_llm_judge()

    final_reward = round(0.60 * (binary_passed / binary_total) + 0.40 * llm_score, 4)

    result = {
        "reward": final_reward,
        "binary": {
            "passed": binary_passed,
            "total": binary_total,
            "score": round(binary_passed / binary_total, 4),
            "detail": {
                f"C{i:02d}": {"passed": ok, "label": lbl}
                for i, (ok, lbl) in enumerate(zip(checks, labels), 1)
            },
        },
        "llm_judge": {
            "score": round(llm_score, 4),
            "weight": 0.40,
            **llm_detail,
        },
        "formula": "0.60 × binary_score + 0.40 × llm_score",
    }

    with open("/logs/verifier/reward.json", "w") as f:
        json.dump(result, f, indent=2)
    with open("/logs/verifier/reward.txt", "w") as f:
        f.write(str(final_reward))

    print(f"Binary: {binary_passed}/{binary_total} ({binary_passed/binary_total:.3f})")
    print(f"LLM judge: {llm_score:.3f}  detail: {llm_detail}")
    print(f"Final reward: {final_reward}")


if __name__ == "__main__":
    main()
